from __future__ import annotations

import secrets
from datetime import date

from django.db.models import Avg, Count, Q
from rest_framework import permissions, viewsets
from rest_framework.decorators import action, api_view, permission_classes
from rest_framework.permissions import AllowAny
from rest_framework.response import Response

from apps.core.models import Empresa

from .models import (
    AccionCAPA, AcuerdoRevision, ActividadSGC, Auditoria, Capacitacion, ParticipanteCapacitacion, ComentarioSGC,
    CompetenciaPerfil, ElementoContexto, Encuesta, Equipo, EvaluacionCompetencia, EvaluacionDetalle, EvaluacionProveedor,
    EvaluacionRequisito, EvidenciaSGC, Hallazgo, IndicadorKPI, MedicionKPI, NoConformidad,
    NotificacionCalidad, Norma, ObjetivoCalidad, ParteInteresada, PerfilPuesto, PoliticaCalidad,
    Proceso, Queja, RequisitoISO, RespuestaEncuesta, RevisionDireccion, Riesgo,
    SalidaNoConforme, SubtareaImplementacion, TareaImplementacion, TareaObjetivo,
    PlantillaDocumento, HitoCertificacion, ProgramaAuditoria, PlantillaChecklist,
    ItemChecklist, RegistroCalidad, GestionCambio, ComunicacionSGC, ConocimientoOrganizacional,
)
from .serializers import (
    AccionCAPASerializer, AcuerdoRevisionSerializer, ActividadSGCSerializer, AuditoriaSerializer,
    CapacitacionSerializer, ParticipanteCapacitacionSerializer,
    ComentarioSGCSerializer, ElementoContextoSerializer, EquipoSerializer,
    EvaluacionCompetenciaSerializer, EvaluacionProveedorSerializer, EvaluacionRequisitoSerializer,
    EvidenciaSGCSerializer, HallazgoSerializer, IndicadorKPISerializer, MedicionKPISerializer,
    MiembroSerializer, NoConformidadSerializer, NotificacionCalidadSerializer, NormaSerializer,
    ObjetivoCalidadSerializer, ParteInteresadaSerializer, PerfilPuestoSerializer,
    PoliticaCalidadSerializer, ProcesoSerializer, QuejaSerializer, RequisitoISOSerializer,
    RevisionDireccionSerializer, RiesgoSerializer, SalidaNoConformeSerializer,
    SubtareaImplementacionSerializer,
    TareaImplementacionSerializer, TareaObjetivoSerializer,
    EncuestaSerializer, EncuestaPublicaSerializer, RespuestaEncuestaSerializer,
    CompetenciaPerfilSerializer, EvaluacionDetalleSerializer,
    PlantillaDocumentoSerializer, HitoCertificacionSerializer, ProgramaAuditoriaSerializer,
    PlantillaChecklistSerializer, ItemChecklistSerializer, RegistroCalidadSerializer,
    GestionCambioSerializer, ComunicacionSGCSerializer, ConocimientoOrganizacionalSerializer,
)
from .servicios import (
    ct_para, extraer_menciones, marcar_leidas, miembros_empresa, nombre_usuario,
    notificar, registrar_actividad, url_objeto,
)


def _empresas(user):
    if user.is_superuser:
        return Empresa.objects.values_list("id", flat=True)
    return user.empresas.filter(activo=True).values_list("empresa_id", flat=True)


class _Scoped(viewsets.ModelViewSet):
    permission_classes = [permissions.IsAuthenticated]
    empresa_field = "empresa"

    def get_queryset(self):
        qs = super().get_queryset()
        return qs.filter(**{f"{self.empresa_field}__in": _empresas(self.request.user)})


class _Colaborativo(_Scoped):
    """ModelViewSet con trazabilidad y notificaciones automáticas.

    Registra cada alta/cambio en la bitácora (ActividadSGC) y notifica al usuario
    asignado cuando se le asigna un registro o cambia su estado. La subclase
    declara `asignacion_field` (FK al User responsable).
    """
    asignacion_field = "responsable_user"

    def _label(self, obj) -> str:
        nombre = obj._meta.verbose_name or obj._meta.model_name
        ref = getattr(obj, "folio", "") or getattr(obj, "titulo", "") or getattr(obj, "nombre", "") or f"#{obj.pk}"
        return f"{nombre} {ref}".strip()

    def _asignado(self, obj):
        return getattr(obj, self.asignacion_field, None)

    def perform_create(self, serializer):
        user = self.request.user
        extra = {}
        if any(f.name == "creado_por" for f in serializer.Meta.model._meta.fields):
            extra["creado_por"] = user
        obj = serializer.save(**extra)
        registrar_actividad(obj, user, "CREO", f"creó {self._label(obj)}")
        asignado = self._asignado(obj)
        if asignado:
            notificar([asignado], obj.empresa, user, "ASIGNACION",
                      f"Te asignaron: {self._label(obj)}", obj=obj)

    def perform_update(self, serializer):
        user = self.request.user
        inst = serializer.instance
        antes_asig = getattr(inst, f"{self.asignacion_field}_id", None)
        antes_estado = getattr(inst, "estado", None)
        obj = serializer.save()
        nuevo_asig = self._asignado(obj)
        if nuevo_asig and getattr(nuevo_asig, "pk", None) != antes_asig:
            registrar_actividad(obj, user, "ASIGNO", f"asignó a {nombre_usuario(nuevo_asig)}")
            notificar([nuevo_asig], obj.empresa, user, "ASIGNACION",
                      f"Te asignaron: {self._label(obj)}", obj=obj)
        nuevo_estado = getattr(obj, "estado", None)
        if antes_estado is not None and nuevo_estado != antes_estado:
            registrar_actividad(obj, user, "CAMBIO_ESTADO",
                                f"cambió el estado a {nuevo_estado}",
                                {"estado": [antes_estado, nuevo_estado]})
            if nuevo_asig:
                notificar([nuevo_asig], obj.empresa, user, "CAMBIO_ESTADO",
                          f"{self._label(obj)}: {antes_estado} → {nuevo_estado}", obj=obj)
        else:
            registrar_actividad(obj, user, "ACTUALIZO", f"actualizó {self._label(obj)}")


# Catálogos de norma/requisitos (no scoped por empresa: son globales).
class NormaViewSet(viewsets.ModelViewSet):
    permission_classes = [permissions.IsAuthenticated]
    queryset = Norma.objects.all()
    serializer_class = NormaSerializer


class RequisitoISOViewSet(viewsets.ModelViewSet):
    permission_classes = [permissions.IsAuthenticated]
    queryset = RequisitoISO.objects.select_related("norma")
    serializer_class = RequisitoISOSerializer
    filterset_fields = ["norma"]


class EvaluacionRequisitoViewSet(_Scoped):
    queryset = EvaluacionRequisito.objects.select_related("requisito")
    serializer_class = EvaluacionRequisitoSerializer
    filterset_fields = ["empresa", "requisito", "requisito__norma", "cumple"]

    def create(self, request, *args, **kwargs):
        """Upsert: si ya existe la evaluación de (empresa, requisito) la actualiza
        en vez de fallar por la restricción única empresa+requisito. Así el
        diagnóstico nunca se rompe aunque el cliente envíe POST en lugar de PATCH."""
        ids = list(_empresas(request.user))
        try:
            emp = int(request.data.get("empresa"))
            req = int(request.data.get("requisito"))
        except (TypeError, ValueError):
            emp = req = None
        if emp and req and emp in ids:
            existente = EvaluacionRequisito.objects.filter(empresa_id=emp, requisito_id=req).first()
            if existente:
                ser = self.get_serializer(existente, data=request.data, partial=True)
                ser.is_valid(raise_exception=True)
                ser.save()
                return Response(ser.data)
        return super().create(request, *args, **kwargs)

    @action(detail=False, methods=["post"])
    def generar_capa(self, request):
        """Crea no conformidades automáticamente desde las brechas (NO/PARCIAL)."""
        ids = list(_empresas(request.user))
        emp_id = request.data.get("empresa") or (ids[0] if ids else None)
        if not emp_id or int(emp_id) not in ids:
            return Response({"detail": "Empresa inválida."}, status=400)
        brechas = (EvaluacionRequisito.objects.filter(empresa_id=emp_id, cumple__in=["NO", "PARCIAL"])
                   .select_related("requisito__norma"))
        creadas = 0
        for ev in brechas:
            desc = f"Brecha {ev.requisito.norma.codigo} {ev.requisito.clausula}: {ev.requisito.titulo}"
            if NoConformidad.objects.filter(empresa_id=emp_id, descripcion=desc).exists():
                continue
            NoConformidad.objects.create(
                empresa_id=emp_id,
                tipo="CORRECTIVA" if ev.cumple == "NO" else "MEJORA", origen="AUDITORIA",
                descripcion=desc, accion=ev.observaciones or "Definir e implementar acción para cerrar la brecha.",
                estado="ABIERTA",
            )
            creadas += 1
        return Response({"creadas": creadas, "brechas": brechas.count()})


class NoConformidadViewSet(_Colaborativo):
    queryset = NoConformidad.objects.select_related("responsable_user", "creado_por")
    serializer_class = NoConformidadSerializer
    filterset_fields = ["empresa", "tipo", "estado", "origen", "responsable_user"]
    search_fields = ["folio", "descripcion"]

    @action(detail=True, methods=["post"])
    def cerrar(self, request, pk=None):
        from datetime import date as _d
        nc = self.get_object()
        antes = nc.estado
        nc.estado = "CERRADA"
        nc.fecha_cierre = _d.today()
        nc.eficacia_verificada = bool(request.data.get("eficacia", True))
        nc.save()
        registrar_actividad(nc, request.user, "CERRO",
                            f"cerró la NC (eficacia {'verificada' if nc.eficacia_verificada else 'pendiente'})",
                            {"estado": [antes, "CERRADA"]})
        if nc.creado_por and nc.creado_por_id != request.user.id:
            notificar([nc.creado_por], nc.empresa, request.user, "CAMBIO_ESTADO",
                      f"Se cerró la NC {nc.folio or nc.id}", obj=nc)
        return Response(NoConformidadSerializer(nc).data)

    @action(detail=True, methods=["post"])
    def aprobar(self, request, pk=None):
        """Aprueba el plan de acción (segregación: no debería ser el ejecutor)."""
        from datetime import date as _d
        nc = self.get_object()
        nc.aprobada_por = request.user
        nc.fecha_aprobacion = _d.today()
        if nc.estado == "ABIERTA":
            nc.estado = "EN_PROCESO"
        nc.save()
        registrar_actividad(nc, request.user, "APROBO", "aprobó el plan de acción")
        destinatarios = [nc.responsable_user, nc.creado_por]
        notificar(destinatarios, nc.empresa, request.user, "APROBACION",
                  f"Plan de acción aprobado: {nc.folio or nc.id}", obj=nc)
        return Response(NoConformidadSerializer(nc).data)

    @action(detail=True, methods=["post"])
    def verificar_eficacia(self, request, pk=None):
        nc = self.get_object()
        eficaz = bool(request.data.get("eficaz", True))
        nc.verificada_por = request.user
        nc.eficacia_verificada = eficaz
        nc.save()
        registrar_actividad(nc, request.user, "VERIFICO",
                            f"verificó la eficacia: {'eficaz' if eficaz else 'no eficaz'}")
        notificar([nc.responsable_user, nc.creado_por], nc.empresa, request.user, "CAMBIO_ESTADO",
                  f"Eficacia {'confirmada' if eficaz else 'no lograda'}: {nc.folio or nc.id}", obj=nc)
        return Response(NoConformidadSerializer(nc).data)


class AuditoriaViewSet(_Colaborativo):
    asignacion_field = "auditor_lider_user"
    queryset = Auditoria.objects.prefetch_related("hallazgos").select_related("auditor_lider_user")
    serializer_class = AuditoriaSerializer
    filterset_fields = ["empresa", "tipo", "estado", "auditor_lider_user"]

    @action(detail=True, methods=["get"])
    def checklist(self, request, pk=None):
        """Devuelve el checklist de ejecución de la auditoría. Si está vacío, lo
        siembra con los requisitos de la norma (o ISO 9001 por defecto)."""
        from .models import RespuestaChecklistAuditoria
        aud = self.get_object()
        existentes = RespuestaChecklistAuditoria.objects.filter(auditoria=aud)
        if not existentes.exists():
            norma = aud.norma or Norma.objects.filter(codigo__icontains="9001").first()
            reqs = RequisitoISO.objects.filter(norma=norma) if norma else RequisitoISO.objects.none()
            for r in reqs:
                RespuestaChecklistAuditoria.objects.create(
                    auditoria=aud, requisito=r, clausula=r.clausula,
                    pregunta=r.titulo, resultado="PENDIENTE")
            existentes = RespuestaChecklistAuditoria.objects.filter(auditoria=aud)
        items = [{
            "id": x.id, "clausula": x.clausula, "pregunta": x.pregunta,
            "resultado": x.resultado, "nota": x.nota,
        } for x in existentes.order_by("id")]
        # Progreso y conteo.
        total = len(items)
        evaluados = sum(1 for x in items if x["resultado"] != "PENDIENTE")
        return Response({
            "auditoria": aud.titulo, "estado": aud.estado, "total": total,
            "evaluados": evaluados,
            "progreso": round(evaluados / total * 100, 1) if total else 0,
            "conforme": sum(1 for x in items if x["resultado"] == "CONFORME"),
            "no_conforme": sum(1 for x in items if x["resultado"] == "NO_CONFORME"),
            "observacion": sum(1 for x in items if x["resultado"] == "OBSERVACION"),
            "items": items,
        })

    @action(detail=True, methods=["post"], url_path="guardar-checklist")
    def guardar_checklist(self, request, pk=None):
        """Guarda la respuesta de un punto del checklist (resultado + nota)."""
        from .models import RespuestaChecklistAuditoria
        aud = self.get_object()
        item_id = request.data.get("id")
        try:
            item = RespuestaChecklistAuditoria.objects.get(id=item_id, auditoria=aud)
        except RespuestaChecklistAuditoria.DoesNotExist:
            return Response({"detail": "Punto no encontrado."}, status=404)
        if "resultado" in request.data:
            item.resultado = request.data["resultado"]
        if "nota" in request.data:
            item.nota = request.data["nota"]
        item.save()
        if aud.estado == "PROGRAMADA":
            aud.estado = "EN_CURSO"
            aud.save(update_fields=["estado"])
        return Response({"ok": True})

    @action(detail=True, methods=["post"], url_path="cerrar-checklist")
    def cerrar_checklist(self, request, pk=None):
        """Cierra la auditoría y genera hallazgos automáticamente desde los
        puntos NO CONFORME y OBSERVACIÓN del checklist."""
        from .models import RespuestaChecklistAuditoria
        aud = self.get_object()
        creados = 0
        for x in RespuestaChecklistAuditoria.objects.filter(
                auditoria=aud, resultado__in=["NO_CONFORME", "OBSERVACION"]):
            # Evita duplicar hallazgos del mismo requisito.
            if Hallazgo.objects.filter(auditoria=aud, requisito=x.requisito).exists():
                continue
            tipo = "NC_MENOR" if x.resultado == "NO_CONFORME" else "OBSERVACION"
            Hallazgo.objects.create(
                auditoria=aud, requisito=x.requisito, tipo=tipo,
                descripcion=f"[{x.clausula}] {x.pregunta}",
                evidencia=x.nota)
            creados += 1
        aud.estado = "CERRADA"
        from datetime import date as _d
        if not aud.fecha_realizada:
            aud.fecha_realizada = _d.today()
        aud.save(update_fields=["estado", "fecha_realizada"])
        registrar_actividad(aud, request.user, "CERRO",
                            f"cerró la auditoría y generó {creados} hallazgo(s)")
        return Response({"hallazgos_creados": creados, "estado": aud.estado})


class HallazgoViewSet(viewsets.ModelViewSet):
    permission_classes = [permissions.IsAuthenticated]
    queryset = Hallazgo.objects.select_related("auditoria")
    serializer_class = HallazgoSerializer
    filterset_fields = ["auditoria", "tipo"]

    def get_queryset(self):
        return super().get_queryset().filter(auditoria__empresa__in=_empresas(self.request.user))

    @action(detail=True, methods=["post"])
    def convertir_nc(self, request, pk=None):
        h = self.get_object()
        if h.no_conformidad_id:
            return Response({"detail": "Ya tiene NC.", "no_conformidad": h.no_conformidad_id})
        nc = NoConformidad.objects.create(
            empresa=h.auditoria.empresa, tipo="CORRECTIVA", origen="AUDITORIA",
            descripcion=f"[{h.auditoria.titulo}] {h.descripcion}", estado="ABIERTA",
        )
        h.no_conformidad = nc
        h.save(update_fields=["no_conformidad"])
        return Response({"no_conformidad": nc.id})


class RiesgoViewSet(_Colaborativo):
    queryset = Riesgo.objects.select_related("responsable_user")
    serializer_class = RiesgoSerializer
    filterset_fields = ["empresa", "estado", "es_oportunidad", "responsable_user"]


class IndicadorKPIViewSet(_Colaborativo):
    queryset = IndicadorKPI.objects.select_related("responsable_user")
    serializer_class = IndicadorKPISerializer
    filterset_fields = ["empresa", "proceso", "responsable_user"]

    @action(detail=True, methods=["get", "post"])
    def mediciones(self, request, pk=None):
        """GET: historial de mediciones del KPI con su tendencia.
        POST: registra una nueva medición (valor + fecha) y actualiza el valor actual."""
        kpi = self.get_object()
        if request.method == "POST":
            from datetime import date as _date
            valor = request.data.get("valor")
            if valor is None or valor == "":
                return Response({"detail": "El valor es obligatorio."}, status=400)
            fecha = request.data.get("fecha") or str(_date.today())
            MedicionKPI.objects.update_or_create(
                kpi=kpi, fecha=fecha,
                defaults={"valor": valor, "nota": request.data.get("nota", ""),
                          "registrado_por": request.user})
            # Actualiza el valor actual del KPI con la última medición.
            ultima = kpi.mediciones.order_by("-fecha").first()
            if ultima:
                kpi.valor_actual = ultima.valor
                kpi.save(update_fields=["valor_actual"])
            registrar_actividad(kpi, request.user, "ACTUALIZO", f"registró medición {valor}")

        mediciones = list(kpi.mediciones.order_by("fecha"))
        meta = float(kpi.meta or 0)
        serie = []
        for m in mediciones:
            v = float(m.valor)
            cumple = (v <= meta) if kpi.sentido == "MENOR" else (v >= meta)
            serie.append({"fecha": str(m.fecha), "valor": v, "meta": meta,
                          "nota": m.nota, "cumple": cumple})
        # Semáforo del valor actual.
        actual = float(kpi.valor_actual or 0)
        if kpi.sentido == "MENOR":
            ratio = (meta / actual) if actual else 1
        else:
            ratio = (actual / meta) if meta else 0
        semaforo = "verde" if ratio >= 1 else "amarillo" if ratio >= 0.85 else "rojo"
        return Response({
            "kpi": kpi.nombre, "unidad": kpi.unidad, "meta": meta, "sentido": kpi.sentido,
            "valor_actual": actual, "cumple": kpi.cumple, "semaforo": semaforo,
            "mediciones": serie,
        })


class CapacitacionViewSet(_Colaborativo):
    queryset = Capacitacion.objects.select_related("responsable_user").prefetch_related(
        "participantes_lista__empleado")
    serializer_class = CapacitacionSerializer
    filterset_fields = ["empresa", "estado", "responsable_user"]

    @action(detail=True, methods=["post"], url_path="reprogramar-recuperacion")
    def reprogramar_recuperacion(self, request, pk=None):
        """Crea una NUEVA edición (curso de recuperación) del mismo curso y mueve
        ahí a los reprobados/ausentes para que lo presenten de nuevo, con un
        intento adicional. Los originales quedan marcados como reprogramados."""
        from datetime import date as _date
        origen = self.get_object()
        # ¿Qué participantes reprogramar? Si se pasan ids, solo esos; si no, todos
        # los reprobados/ausentes aún no reprogramados de esta capacitación.
        ids_part = request.data.get("participantes")
        pendientes = origen.participantes_lista.filter(
            estado__in=["NO_APROBADO", "AUSENTE"], reprogramado=False)
        if ids_part:
            pendientes = pendientes.filter(id__in=ids_part)
        pendientes = list(pendientes.select_related("empleado"))
        if not pendientes:
            return Response({"detail": "No hay reprobados por reprogramar."}, status=400)

        fecha = request.data.get("fecha") or None

        # Crea la nueva edición copiando los datos del curso.
        nueva = Capacitacion.objects.create(
            empresa=origen.empresa,
            curso=origen.curso,
            descripcion=origen.descripcion,
            instructor=origen.instructor,
            responsable_user=origen.responsable_user,
            fecha=fecha,
            vigencia_meses=origen.vigencia_meses,
            requiere_calificacion=origen.requiere_calificacion,
            calificacion_minima=origen.calificacion_minima,
            estado="PROGRAMADA",
        )
        movidos = 0
        for p in pendientes:
            ParticipanteCapacitacion.objects.create(
                capacitacion=nueva,
                empleado=p.empleado,
                nombre=p.nombre,
                estado="INSCRITO",
                intentos=(p.intentos or 1) + 1,
            )
            if p.empleado_id:
                nueva.empleados.add(p.empleado)
            p.reprogramado = True
            p.save(update_fields=["reprogramado"])
            movidos += 1

        registrar_actividad(origen, request.user, "CREO",
                            f"reprogramó recuperación de {movidos} reprobado(s) a una nueva edición")
        return Response({"nueva_capacitacion": nueva.id, "movidos": movidos,
                         "curso": nueva.curso}, status=201)

    @action(detail=False, methods=["get"], url_path="pendientes-recapacitacion")
    def pendientes_recapacitacion(self, request):
        """Lista de participantes que NO aprobaron y deben volver a presentar el
        curso (estado NO_APROBADO o AUSENTE), aún no reprogramados a otra edición."""
        ids = self._scoped_ids(request)
        qs = (ParticipanteCapacitacion.objects
              .filter(capacitacion__empresa_id__in=ids, estado__in=["NO_APROBADO", "AUSENTE"],
                      reprogramado=False)
              .select_related("capacitacion", "empleado")
              .order_by("capacitacion__curso", "id"))
        out = []
        for p in qs:
            out.append({
                "id": p.id,
                "nombre": p.nombre_display,
                "empleado": p.empleado_id,
                "empleado_numero": p.empleado.numero_empleado if p.empleado_id else None,
                "curso": p.capacitacion.curso,
                "capacitacion_id": p.capacitacion_id,
                "estado": p.estado,
                "estado_display": p.get_estado_display(),
                "calificacion": float(p.calificacion) if p.calificacion is not None else None,
                "calificacion_minima": float(p.capacitacion.calificacion_minima),
                "intentos": p.intentos,
                "fecha": p.capacitacion.fecha,
            })
        return Response({"total": len(out), "resultados": out})

    def _scoped_ids(self, request):
        ids = list(_empresas(request.user))
        emp = request.query_params.get("empresa")
        return [int(emp)] if emp and int(emp) in ids else ids


class ParticipanteCapacitacionViewSet(viewsets.ModelViewSet):
    """Participantes individuales de una capacitación (calificación, evidencia,
    estado por persona). Soporta empleados de RH o participantes externos."""
    permission_classes = [permissions.IsAuthenticated]
    serializer_class = ParticipanteCapacitacionSerializer
    filterset_fields = ["capacitacion", "empleado", "estado"]

    def get_queryset(self):
        return (ParticipanteCapacitacion.objects
                .filter(capacitacion__empresa__in=_empresas(self.request.user))
                .select_related("empleado", "capacitacion"))

    def _aplicar_calificacion(self, obj):
        """Si el curso requiere calificación y hay nota, fija aprobado/no aprobado."""
        cap = obj.capacitacion
        if obj.calificacion is not None and cap.requiere_calificacion:
            obj.evaluar_estado(cap.calificacion_minima)
            obj.save(update_fields=["estado"])

    def perform_create(self, serializer):
        obj = serializer.save()
        self._aplicar_calificacion(obj)
        # Si el participante es un empleado de RH, vincúlalo también al M2M
        # para que la capacitación aparezca en su portal del empleado.
        if obj.empleado_id:
            obj.capacitacion.empleados.add(obj.empleado)

    def perform_update(self, serializer):
        obj = serializer.save()
        self._aplicar_calificacion(obj)

    @action(detail=True, methods=["post"], url_path="evidencia")
    def subir_evidencia(self, request, pk=None):
        """Sube la constancia/evidencia individual del participante (multipart)."""
        p = self.get_object()
        archivo = request.FILES.get("archivo")
        if not archivo:
            return Response({"detail": "Adjunta el archivo en el campo 'archivo'."}, status=400)
        p.evidencia = archivo
        p.evidencia_nombre = getattr(archivo, "name", "")[:200]
        p.save(update_fields=["evidencia", "evidencia_nombre"])
        return Response(ParticipanteCapacitacionSerializer(p, context={"request": request}).data)

    @action(detail=True, methods=["post"], url_path="reinscribir")
    def reinscribir(self, request, pk=None):
        """Reinscribe a un participante que no aprobó: borra su calificación,
        lo regresa a INSCRITO e incrementa el número de intentos. El curso debe
        repetirse hasta aprobar."""
        p = self.get_object()
        p.intentos = (p.intentos or 1) + 1
        p.calificacion = None
        p.estado = "INSCRITO"
        p.save(update_fields=["intentos", "calificacion", "estado"])
        return Response(ParticipanteCapacitacionSerializer(p, context={"request": request}).data)

    @action(detail=True, methods=["get"], url_path="constancia")
    def constancia(self, request, pk=None):
        """Genera la constancia PDF del participante (si aprobó o el curso es solo
        asistencia). Reconocimiento oficial de la capacitación recibida."""
        from django.http import HttpResponse
        p = self.get_object()
        cap = p.capacitacion
        # Solo emite constancia a quien aprobó, asistió, o cursos sin calificación.
        if cap.requiere_calificacion and p.estado not in ("APROBADO", "ASISTIO"):
            return Response(
                {"detail": "La constancia solo se emite a participantes aprobados."}, status=400)

        import io
        from reportlab.lib import colors
        from reportlab.lib.pagesizes import landscape, letter
        from reportlab.lib.units import cm
        from reportlab.pdfgen import canvas as _canvas

        emp = cap.empresa
        empresa_nombre = emp.nombre_comercial or emp.razon_social or "Empresa"
        buf = io.BytesIO()
        W, H = landscape(letter)
        c = _canvas.Canvas(buf, pagesize=landscape(letter))

        # Marco decorativo
        c.setStrokeColor(colors.HexColor("#1A73E8"))
        c.setLineWidth(3)
        c.rect(1.2 * cm, 1.2 * cm, W - 2.4 * cm, H - 2.4 * cm)
        c.setStrokeColor(colors.HexColor("#14B8A6"))
        c.setLineWidth(1)
        c.rect(1.5 * cm, 1.5 * cm, W - 3.0 * cm, H - 3.0 * cm)

        cx = W / 2
        c.setFillColor(colors.HexColor("#0F172A"))
        c.setFont("Helvetica-Bold", 12)
        c.drawCentredString(cx, H - 2.6 * cm, empresa_nombre.upper())
        c.setFont("Helvetica", 9)
        c.setFillColor(colors.HexColor("#64748B"))
        c.drawCentredString(cx, H - 3.1 * cm, "Sistema de Gestión de Calidad ISO 9001:2015")

        c.setFont("Helvetica-Bold", 30)
        c.setFillColor(colors.HexColor("#1A73E8"))
        c.drawCentredString(cx, H - 5.2 * cm, "CONSTANCIA")
        c.setFont("Helvetica", 12)
        c.setFillColor(colors.HexColor("#334155"))
        c.drawCentredString(cx, H - 6.1 * cm, "Se otorga la presente constancia a:")

        c.setFont("Helvetica-Bold", 22)
        c.setFillColor(colors.HexColor("#0F172A"))
        c.drawCentredString(cx, H - 7.6 * cm, p.nombre_display)
        # Subrayado del nombre
        c.setStrokeColor(colors.HexColor("#14B8A6"))
        c.setLineWidth(1)
        nw = c.stringWidth(p.nombre_display, "Helvetica-Bold", 22)
        c.line(cx - nw / 2 - 20, H - 7.9 * cm, cx + nw / 2 + 20, H - 7.9 * cm)

        c.setFont("Helvetica", 12)
        c.setFillColor(colors.HexColor("#334155"))
        c.drawCentredString(cx, H - 8.8 * cm, "por haber concluido satisfactoriamente el curso:")
        c.setFont("Helvetica-Bold", 16)
        c.setFillColor(colors.HexColor("#1A73E8"))
        c.drawCentredString(cx, H - 9.7 * cm, f"“{cap.curso}”")

        # Detalles
        c.setFont("Helvetica", 11)
        c.setFillColor(colors.HexColor("#475569"))
        detalles = []
        if cap.instructor:
            detalles.append(f"Instructor: {cap.instructor}")
        if cap.fecha:
            detalles.append(f"Fecha: {cap.fecha.strftime('%d/%m/%Y')}")
        if cap.requiere_calificacion and p.calificacion is not None:
            detalles.append(f"Calificación: {p.calificacion}")
        if detalles:
            c.drawCentredString(cx, H - 10.7 * cm, "   ·   ".join(detalles))
        if cap.fecha_vencimiento:
            c.setFillColor(colors.HexColor("#94A3B8"))
            c.setFont("Helvetica-Oblique", 9)
            c.drawCentredString(cx, H - 11.3 * cm, f"Vigencia hasta: {cap.fecha_vencimiento.strftime('%d/%m/%Y')}")

        # Firmas
        y_firma = 2.6 * cm
        for fx, rol in [(cx - 6 * cm, "Instructor"), (cx + 6 * cm, "Responsable de Calidad")]:
            c.setStrokeColor(colors.HexColor("#94A3B8"))
            c.setLineWidth(0.7)
            c.line(fx - 3.5 * cm, y_firma, fx + 3.5 * cm, y_firma)
            c.setFont("Helvetica", 9)
            c.setFillColor(colors.HexColor("#64748B"))
            c.drawCentredString(fx, y_firma - 0.5 * cm, rol)

        from datetime import date as _date
        c.setFont("Helvetica", 8)
        c.setFillColor(colors.HexColor("#CBD5E1"))
        c.drawCentredString(cx, 1.8 * cm, f"Emitida el {_date.today().strftime('%d/%m/%Y')} · Folio CONST-{cap.id}-{p.id}")

        c.showPage()
        c.save()
        pdf = buf.getvalue()
        resp = HttpResponse(pdf, content_type="application/pdf")
        nombre_arch = f"constancia-{p.nombre_display.replace(' ', '_')}-{cap.id}.pdf"
        resp["Content-Disposition"] = f'inline; filename="{nombre_arch}"'
        return resp


class EquipoViewSet(_Colaborativo):
    queryset = Equipo.objects.select_related("responsable_user")
    serializer_class = EquipoSerializer
    filterset_fields = ["empresa", "activo", "responsable_user"]


class EvaluacionProveedorViewSet(_Colaborativo):
    queryset = EvaluacionProveedor.objects.select_related("proveedor", "responsable_user")
    serializer_class = EvaluacionProveedorSerializer
    filterset_fields = ["empresa", "proveedor", "estado", "responsable_user"]

    def _label(self, obj):
        return f"evaluación de {obj.proveedor.razon_social} ({obj.periodo})"

    def perform_create(self, serializer):
        super().perform_create(serializer)  # trazabilidad + notifica al responsable
        ev = serializer.instance
        # Alerta automática: si el proveedor cae a clasificación D (<70),
        # notifica a la dirección/calidad de la empresa.
        if float(ev.puntaje) < 70:
            from django.contrib.auth import get_user_model
            from apps.core.models import UsuarioEmpresa
            uids = UsuarioEmpresa.objects.filter(
                empresa=ev.empresa, activo=True, rol__in=["OWNER", "STAFF", "MANAGER"]
            ).values_list("user_id", flat=True)
            dest = list(get_user_model().objects.filter(id__in=uids))
            notificar(dest, ev.empresa, self.request.user, "VENCIMIENTO",
                      f"Proveedor crítico: {ev.proveedor.razon_social}",
                      mensaje=f"Cayó a clasificación D (puntaje {ev.puntaje}) en {ev.periodo}. Requiere plan de mejora.",
                      url="/sgc/evaluacion-proveedores")


class QuejaViewSet(_Colaborativo):
    queryset = Queja.objects.select_related("responsable_user")
    serializer_class = QuejaSerializer
    filterset_fields = ["empresa", "tipo", "estado", "responsable_user"]

    @action(detail=True, methods=["post"])
    def escalar(self, request, pk=None):
        q = self.get_object()
        if q.no_conformidad_id:
            return Response({"detail": "Ya escalada.", "no_conformidad": q.no_conformidad_id})
        nc = NoConformidad.objects.create(
            empresa=q.empresa, tipo="CORRECTIVA", origen="QUEJA",
            descripcion=f"[Queja {q.folio or q.id}] {q.descripcion}", estado="ABIERTA",
            responsable_user=q.responsable_user,
        )
        q.no_conformidad = nc
        q.estado = "EN_PROCESO"
        q.save(update_fields=["no_conformidad", "estado"])
        registrar_actividad(q, request.user, "CAMBIO_ESTADO",
                            f"escaló la queja a la no conformidad {nc.folio}")
        return Response({"no_conformidad": nc.id})


class SalidaNoConformeViewSet(_Colaborativo):
    """Control de salidas no conformes (ISO 9001 8.7)."""
    queryset = SalidaNoConforme.objects.select_related("responsable_user", "proceso_ref", "no_conformidad")
    serializer_class = SalidaNoConformeSerializer
    filterset_fields = ["empresa", "estado", "origen", "disposicion", "responsable_user"]

    def _label(self, obj):
        return f"salida no conforme {obj.folio}"

    @action(detail=True, methods=["post"])
    def escalar(self, request, pk=None):
        """Eleva la salida no conforme a una No Conformidad del sistema (10.2)
        cuando el problema es recurrente o grave."""
        snc = self.get_object()
        if snc.no_conformidad_id:
            return Response({"detail": "Ya escalada.", "no_conformidad": snc.no_conformidad_id})
        nc = NoConformidad.objects.create(
            empresa=snc.empresa, tipo="CORRECTIVA", origen="PROCESO",
            descripcion=f"[Salida no conforme {snc.folio}] {snc.descripcion}",
            estado="ABIERTA", responsable_user=snc.responsable_user,
        )
        snc.no_conformidad = nc
        snc.save(update_fields=["no_conformidad"])
        registrar_actividad(snc, request.user, "CAMBIO_ESTADO",
                            f"escaló a la no conformidad {nc.folio}")
        return Response({"no_conformidad": nc.id})


class PoliticaCalidadViewSet(_Scoped):
    queryset = PoliticaCalidad.objects.all()
    serializer_class = PoliticaCalidadSerializer
    filterset_fields = ["empresa"]


class ObjetivoCalidadViewSet(_Colaborativo):
    queryset = ObjetivoCalidad.objects.select_related("responsable_user")
    serializer_class = ObjetivoCalidadSerializer
    filterset_fields = ["empresa", "estado", "responsable_user"]


class ElementoContextoViewSet(_Colaborativo):
    queryset = ElementoContexto.objects.select_related("responsable_user")
    serializer_class = ElementoContextoSerializer
    filterset_fields = ["empresa", "tipo", "responsable_user"]

    def _label(self, obj):
        return f"factor FODA: {obj.descripcion[:40]}"


class ParteInteresadaViewSet(_Scoped):
    queryset = ParteInteresada.objects.all()
    serializer_class = ParteInteresadaSerializer
    filterset_fields = ["empresa", "tipo"]


def _snapshot_sgc(ids):
    """Foto compacta del estado del SGC para la revisión por la dirección (9.3.2)."""
    nc = NoConformidad.objects.filter(empresa_id__in=ids)
    aud = Auditoria.objects.filter(empresa_id__in=ids)
    kpis = list(IndicadorKPI.objects.filter(empresa_id__in=ids))
    riesgos = list(Riesgo.objects.filter(empresa_id__in=ids))
    quejas = Queja.objects.filter(empresa_id__in=ids)
    sat = quejas.exclude(satisfaccion__isnull=True).aggregate(p=Avg("satisfaccion"))["p"]
    evals = EvaluacionRequisito.objects.filter(empresa_id__in=ids).exclude(cumple="NA")
    puntos = {"NO": 0, "PARCIAL": 50, "SI": 100}
    evd = evals.count()
    cumplimiento = round(sum(puntos.get(e.cumple, 0) for e in evals) / evd, 1) if evd else 0
    impl = TareaImplementacion.objects.filter(empresa_id__in=ids)
    objs = ObjetivoCalidad.objects.filter(empresa_id__in=ids)
    return {
        "cumplimiento_iso": cumplimiento,
        "nc_abiertas": nc.filter(estado__in=["ABIERTA", "EN_PROCESO"]).count(),
        "nc_cerradas": nc.filter(estado="CERRADA").count(),
        "nc_total": nc.count(),
        "auditorias_cerradas": aud.filter(estado="CERRADA").count(),
        "auditorias_programadas": aud.filter(estado="PROGRAMADA").count(),
        "kpis_en_meta": sum(1 for k in kpis if k.cumple),
        "kpis_total": len(kpis),
        "riesgos_altos": sum(1 for r in riesgos if r.severidad in ("ALTO", "CRITICO")),
        "riesgos_total": len(riesgos),
        "satisfaccion": round(sat, 1) if sat else None,
        "quejas_abiertas": quejas.filter(estado__in=["ABIERTA", "EN_PROCESO"]).count(),
        "implementacion_pct": round(impl.filter(columna="VERIFICADO").count() / impl.count() * 100, 1) if impl.count() else 0,
        "objetivos_logrados": objs.filter(estado="LOGRADO").count(),
        "objetivos_total": objs.count(),
        "fecha": str(date.today()),
    }


class RevisionDireccionViewSet(_Scoped):
    queryset = RevisionDireccion.objects.prefetch_related("acuerdos_items").select_related("creado_por")
    serializer_class = RevisionDireccionSerializer
    filterset_fields = ["empresa"]

    def perform_create(self, serializer):
        rev = serializer.save(creado_por=self.request.user)
        rev.metricas = _snapshot_sgc([rev.empresa_id])
        rev.save(update_fields=["metricas"])

    @action(detail=False, methods=["get"])
    def snapshot(self, request):
        """Estado actual del SGC en vivo (para mostrar al armar la revisión)."""
        ids = list(_empresas(request.user))
        emp = request.query_params.get("empresa")
        if emp and int(emp) in ids:
            ids = [int(emp)]
        return Response(_snapshot_sgc(ids))

    @action(detail=False, methods=["get"])
    def datos_sugeridos(self, request):
        """Pre-llena las entradas de la revisión con el estado actual del SGC."""
        ids = list(_empresas(request.user))
        emp = request.query_params.get("empresa")
        if emp and int(emp) in ids:
            ids = [int(emp)]
        from django.db.models import Avg
        nc = NoConformidad.objects.filter(empresa_id__in=ids)
        aud = Auditoria.objects.filter(empresa_id__in=ids)
        kpis = list(IndicadorKPI.objects.filter(empresa_id__in=ids))
        riesgos = list(Riesgo.objects.filter(empresa_id__in=ids))
        quejas = Queja.objects.filter(empresa_id__in=ids)
        sat = quejas.exclude(satisfaccion__isnull=True).aggregate(p=Avg("satisfaccion"))["p"]
        texto = (
            f"DESEMPEÑO DEL SGC:\n"
            f"• No conformidades: {nc.count()} (abiertas: {nc.filter(estado__in=['ABIERTA','EN_PROCESO']).count()}).\n"
            f"• Auditorías realizadas: {aud.filter(estado='CERRADA').count()}; programadas: {aud.filter(estado='PROGRAMADA').count()}.\n"
            f"• KPIs en meta: {sum(1 for k in kpis if k.cumple)}/{len(kpis)}.\n"
            f"• Riesgos críticos/altos: {sum(1 for r in riesgos if r.severidad in ('CRITICO','ALTO'))} de {len(riesgos)}.\n"
            f"• Satisfacción del cliente (prom.): {round(sat,1) if sat else 'N/D'}/5.\n"
            f"• Quejas abiertas: {quejas.filter(estado__in=['ABIERTA','EN_PROCESO']).count()}."
        )
        return Response({"entradas": texto})


class PerfilPuestoViewSet(_Scoped):
    queryset = PerfilPuesto.objects.all()
    serializer_class = PerfilPuestoSerializer
    filterset_fields = ["empresa", "area"]


class EvaluacionCompetenciaViewSet(_Scoped):
    queryset = EvaluacionCompetencia.objects.select_related("perfil").prefetch_related("detalles")
    serializer_class = EvaluacionCompetenciaSerializer
    filterset_fields = ["empresa", "estado", "perfil"]

    @action(detail=True, methods=["post"])
    def sembrar_detalles(self, request, pk=None):
        """Copia las competencias del perfil como detalles a evaluar (nivel 0)."""
        ev = self.get_object()
        if not ev.perfil:
            return Response({"detail": "La evaluación no tiene perfil."}, status=400)
        ev.detalles.all().delete()
        for c in ev.perfil.competencias.all():
            EvaluacionDetalle.objects.create(
                evaluacion=ev, competencia=c, nombre=c.nombre, tipo=c.tipo,
                nivel_requerido=c.nivel_requerido, nivel_actual=0,
                obligatoria=c.obligatoria, orden=c.orden,
            )
        ev.recomputar()
        return Response(EvaluacionCompetenciaSerializer(ev).data)

    @action(detail=True, methods=["post"])
    def crear_capacitacion(self, request, pk=None):
        """Genera un plan de capacitación a partir de las brechas detectadas."""
        from datetime import date as _d, timedelta
        ev = self.get_object()
        brechas = [d for d in ev.detalles.all() if d.nivel_actual < d.nivel_requerido]
        if not brechas:
            return Response({"detail": "Sin brechas que cerrar."}, status=400)
        temas = ", ".join(d.nombre for d in brechas)
        cap = Capacitacion.objects.create(
            empresa=ev.empresa,
            curso=f"Plan de desarrollo · {ev.persona}",
            descripcion=f"Cerrar brechas de competencia de {ev.persona}: {temas}.",
            participantes=ev.persona, estado="PROGRAMADA",
            responsable_user=ev.persona_user,
        )
        ev.capacitacion_ref = cap
        ev.plan_desarrollo = f"Capacitación en: {temas}."
        ev.fecha_reevaluacion = _d.today() + timedelta(days=90)
        ev.save(update_fields=["capacitacion_ref", "plan_desarrollo", "fecha_reevaluacion"])
        return Response({"capacitacion": cap.id, "evaluacion": EvaluacionCompetenciaSerializer(ev).data})

    @action(detail=False, methods=["get"])
    def brechas(self, request):
        """Dashboard de brechas: por competencia (cuántas personas no la cubren)
        y por persona (cuántas brechas tiene)."""
        ids = list(_empresas(request.user))
        emp = request.query_params.get("empresa")
        if emp and int(emp) in ids:
            ids = [int(emp)]
        dets = (EvaluacionDetalle.objects
                .filter(evaluacion__empresa_id__in=ids)
                .select_related("evaluacion"))
        por_comp, por_persona = {}, {}
        total_brechas = 0
        for d in dets:
            if d.nivel_actual < d.nivel_requerido:
                total_brechas += 1
                por_comp[d.nombre] = por_comp.get(d.nombre, 0) + 1
                p = d.evaluacion.persona
                por_persona[p] = por_persona.get(p, 0) + 1
        return Response({
            "total_brechas": total_brechas,
            "por_competencia": sorted(
                [{"nombre": k, "personas": v} for k, v in por_comp.items()],
                key=lambda x: -x["personas"])[:15],
            "por_persona": sorted(
                [{"persona": k, "brechas": v} for k, v in por_persona.items()],
                key=lambda x: -x["brechas"])[:15],
        })


class DashboardSGCViewSet(viewsets.ViewSet):
    """Dashboard ejecutivo del SGC: cumplimiento ISO, NC, auditorías, riesgos…"""
    permission_classes = [permissions.IsAuthenticated]

    def _emp(self, request):
        emp = request.query_params.get("empresa")
        ids = list(_empresas(request.user))
        return [int(emp)] if emp and int(emp) in ids else list(ids)

    @action(detail=False, methods=["get"])
    def resumen(self, request):
        ids = self._emp(request)
        hoy = date.today()

        # % cumplimiento ISO (promedio de puntaje de las evaluaciones, ignora NA).
        evals = EvaluacionRequisito.objects.filter(empresa_id__in=ids).exclude(cumple="NA")
        total_req = RequisitoISO.objects.count()
        puntos = {"NO": 0, "PARCIAL": 50, "SI": 100}
        suma = sum(puntos.get(e.cumple, 0) for e in evals)
        evaluados = evals.count()
        cumplimiento = round(suma / evaluados, 1) if evaluados else 0

        nc = NoConformidad.objects.filter(empresa_id__in=ids)
        aud = Auditoria.objects.filter(empresa_id__in=ids)
        riesgos = list(Riesgo.objects.filter(empresa_id__in=ids))
        kpis = list(IndicadorKPI.objects.filter(empresa_id__in=ids))
        equipos = Equipo.objects.filter(empresa_id__in=ids, activo=True)
        caps = Capacitacion.objects.filter(empresa_id__in=ids)
        quejas = Queja.objects.filter(empresa_id__in=ids)

        riesgos_criticos = sum(1 for r in riesgos if r.severidad == "CRITICO")
        riesgos_altos = sum(1 for r in riesgos if r.severidad == "ALTO")
        kpis_cumplen = sum(1 for k in kpis if k.cumple)
        equipos_vencidos = equipos.filter(
            requiere_calibracion=True, fecha_proxima_calibracion__lt=hoy).count()
        caps_pendientes = caps.filter(estado__in=["PROGRAMADA", "VENCIDA"]).count()

        # Avance global: implementación, objetivos y NC resueltas.
        impl = TareaImplementacion.objects.filter(empresa_id__in=ids)
        impl_total = impl.count()
        impl_verif = impl.filter(columna="VERIFICADO").count()
        objs = ObjetivoCalidad.objects.filter(empresa_id__in=ids)
        obj_total = objs.count()
        obj_logrados = objs.filter(estado="LOGRADO").count()
        obj_avance = round(sum(o.avance for o in objs) / obj_total, 1) if obj_total else 0
        nc_total = nc.count()
        nc_cerradas = nc.filter(estado="CERRADA").count()
        kpis_pct = round(kpis_cumplen / len(kpis) * 100, 1) if kpis else 0

        return Response({
            "cumplimiento_iso": cumplimiento,
            "implementacion": {
                "total": impl_total, "verificadas": impl_verif,
                "progreso": round(impl_verif / impl_total * 100, 1) if impl_total else 0,
            },
            "objetivos": {
                "total": obj_total, "logrados": obj_logrados, "avance_promedio": obj_avance,
                "progreso": round(obj_logrados / obj_total * 100, 1) if obj_total else 0,
            },
            "nc_cerradas": nc_cerradas,
            "nc_resueltas_pct": round(nc_cerradas / nc_total * 100, 1) if nc_total else 0,
            "kpis_meta_pct": kpis_pct,
            "requisitos_total": total_req,
            "requisitos_evaluados": evaluados,
            "no_conformidades": {
                "abiertas": nc.filter(estado__in=["ABIERTA", "EN_PROCESO"]).count(),
                "vencidas": nc.filter(estado__in=["ABIERTA", "EN_PROCESO"], fecha_compromiso__lt=hoy).count(),
                "total": nc.count(),
            },
            "auditorias": {
                "programadas": aud.filter(estado="PROGRAMADA").count(),
                "en_curso": aud.filter(estado="EN_CURSO").count(),
            },
            "riesgos": {"criticos": riesgos_criticos, "altos": riesgos_altos, "total": len(riesgos)},
            "kpis": {"cumplen": kpis_cumplen, "total": len(kpis)},
            "equipos_calibracion_vencida": equipos_vencidos,
            "capacitaciones_pendientes": caps_pendientes,
            "quejas_abiertas": quejas.filter(estado__in=["ABIERTA", "EN_PROCESO"]).count(),
        })

    @action(detail=False, methods=["get"])
    def agenda(self, request):
        """Agenda de cumplimiento: todos los vencimientos del SGC para mantenerlo
        vivo (auditorías, calibraciones, capacitaciones, NC, objetivos, revisión
        por la dirección, acuerdos, tareas y acciones CAPA)."""
        ids = self._emp(request)
        hoy = date.today()
        ev = []

        def add(tipo, titulo, fecha, url):
            if not fecha:
                return
            dias = (fecha - hoy).days
            ev.append({
                "tipo": tipo, "titulo": titulo, "fecha": str(fecha), "dias": dias,
                "estado": "vencido" if dias < 0 else "proximo" if dias <= 30 else "futuro",
                "url": url,
            })

        for a in Auditoria.objects.filter(empresa_id__in=ids, estado__in=["PROGRAMADA", "EN_CURSO"]):
            add("Auditoría", a.titulo, a.fecha_programada, "/sgc/auditorias")
        for e in Equipo.objects.filter(empresa_id__in=ids, requiere_calibracion=True, activo=True):
            add("Calibración", f"{e.codigo} · {e.nombre}", e.fecha_proxima_calibracion, "/sgc/equipos")
        for c in Capacitacion.objects.filter(empresa_id__in=ids).exclude(fecha_vencimiento=None):
            add("Capacitación", c.curso, c.fecha_vencimiento, "/sgc/capacitacion")
        for n in NoConformidad.objects.filter(empresa_id__in=ids, estado__in=["ABIERTA", "EN_PROCESO"]).exclude(fecha_compromiso=None):
            add("No conformidad", f"{n.folio or 'NC'} · {n.descripcion[:50]}", n.fecha_compromiso, "/sgc/no-conformidades")
        for o in ObjetivoCalidad.objects.filter(empresa_id__in=ids, estado="EN_CURSO").exclude(fecha_limite=None):
            add("Objetivo", o.objetivo[:50], o.fecha_limite, "/sgc/politica")
        for r in RevisionDireccion.objects.filter(empresa_id__in=ids).exclude(proxima_fecha=None):
            add("Revisión dirección", f"Revisión {r.periodo or ''}".strip(), r.proxima_fecha, "/sgc/revision-direccion")
        for ac in AcuerdoRevision.objects.filter(revision__empresa_id__in=ids).exclude(estado="CERRADO").exclude(fecha_compromiso=None):
            add("Acuerdo dirección", ac.descripcion[:50], ac.fecha_compromiso, "/sgc/revision-direccion")
        for t in TareaImplementacion.objects.filter(empresa_id__in=ids).exclude(columna="VERIFICADO").exclude(fecha_limite=None):
            add("Implementación", t.titulo[:50], t.fecha_limite, "/sgc/implementacion")
        for acc in AccionCAPA.objects.filter(no_conformidad__empresa_id__in=ids).exclude(estado="VERIFICADA").exclude(fecha_compromiso=None):
            add("Acción CAPA", acc.descripcion[:50], acc.fecha_compromiso, "/sgc/no-conformidades")

        ev.sort(key=lambda x: x["fecha"])
        resumen = {
            "vencidos": sum(1 for e in ev if e["dias"] < 0),
            "semana": sum(1 for e in ev if 0 <= e["dias"] <= 7),
            "mes": sum(1 for e in ev if 0 <= e["dias"] <= 30),
            "total": len(ev),
        }
        return Response({"eventos": ev, "resumen": resumen})

    @action(detail=False, methods=["get"])
    def panorama(self, request):
        """Panorama ejecutivo del SGC para el centro de mando: cumplimiento por
        cláusula ISO, índice de madurez multidimensional y carga del equipo."""
        ids = self._emp(request)
        hoy = date.today()
        puntos = {"NO": 0, "PARCIAL": 50, "SI": 100}

        # ── Cumplimiento por capítulo ISO (4 a 10) ──────────────────────────
        evals = list(EvaluacionRequisito.objects.filter(empresa_id__in=ids)
                     .exclude(cumple="NA").select_related("requisito"))
        por_cap: dict[str, dict] = {}
        for e in evals:
            cap = (e.requisito.clausula or "0").split(".")[0]
            d = por_cap.setdefault(cap, {"suma": 0, "n": 0})
            d["suma"] += puntos.get(e.cumple, 0)
            d["n"] += 1
        CAP_NOMBRE = {
            "4": "Contexto", "5": "Liderazgo", "6": "Planificación",
            "7": "Apoyo / Recursos", "8": "Operación",
            "9": "Evaluación", "10": "Mejora",
        }
        capitulos = []
        for cap in ["4", "5", "6", "7", "8", "9", "10"]:
            d = por_cap.get(cap)
            pct = round(d["suma"] / d["n"], 1) if d and d["n"] else 0
            capitulos.append({
                "capitulo": cap, "nombre": CAP_NOMBRE.get(cap, cap),
                "cumplimiento": pct, "evaluados": d["n"] if d else 0,
            })
        evaluados = len(evals)
        cumplimiento = round(sum(puntos.get(e.cumple, 0) for e in evals) / evaluados, 1) if evaluados else 0

        # ── Índice de madurez (radar multidimensional) ──────────────────────
        impl = TareaImplementacion.objects.filter(empresa_id__in=ids)
        impl_total = impl.count()
        impl_pct = round(impl.filter(columna="VERIFICADO").count() / impl_total * 100, 1) if impl_total else 0
        objs = ObjetivoCalidad.objects.filter(empresa_id__in=ids)
        obj_pct = round(sum(o.avance for o in objs) / objs.count(), 1) if objs.count() else 0
        kpis = list(IndicadorKPI.objects.filter(empresa_id__in=ids))
        kpis_pct = round(sum(1 for k in kpis if k.cumple) / len(kpis) * 100, 1) if kpis else 0
        nc = NoConformidad.objects.filter(empresa_id__in=ids)
        nc_total = nc.count()
        nc_pct = round(nc.filter(estado="CERRADA").count() / nc_total * 100, 1) if nc_total else 0
        riesgos = list(Riesgo.objects.filter(empresa_id__in=ids))
        riesgos_controlados = sum(1 for r in riesgos if r.estado in ("CONTROLADO", "ACEPTADO"))
        riesgo_pct = round(riesgos_controlados / len(riesgos) * 100, 1) if riesgos else 0
        dimensiones = [
            {"dim": "Cumplimiento ISO", "valor": cumplimiento},
            {"dim": "Implementación", "valor": impl_pct},
            {"dim": "Objetivos", "valor": obj_pct},
            {"dim": "KPIs en meta", "valor": kpis_pct},
            {"dim": "NC resueltas", "valor": nc_pct},
            {"dim": "Riesgos controlados", "valor": riesgo_pct},
        ]
        madurez = round(sum(d["valor"] for d in dimensiones) / len(dimensiones), 1)
        if madurez >= 85:
            nivel = "Consolidado"
        elif madurez >= 65:
            nivel = "Maduro"
        elif madurez >= 40:
            nivel = "En desarrollo"
        else:
            nivel = "Inicial"

        # ── Carga del equipo (quién tiene qué asignado y abierto) ───────────
        emp_id = ids[0] if ids else None
        empresa = Empresa.objects.filter(id=emp_id).first() if emp_id else None
        carga = []
        if empresa:
            fuentes = [
                (NoConformidad, "responsable_user", {"estado__in": ["ABIERTA", "EN_PROCESO"]}),
                (Riesgo, "responsable_user", {"estado__in": ["IDENTIFICADO", "EN_TRATAMIENTO"]}),
                (Auditoria, "auditor_lider_user", {"estado__in": ["PROGRAMADA", "EN_CURSO"]}),
                (Capacitacion, "responsable_user", {"estado__in": ["PROGRAMADA"]}),
                (Equipo, "responsable_user", {"activo": True}),
                (ObjetivoCalidad, "responsable_user", {"estado": "EN_CURSO"}),
                (TareaImplementacion, "responsable_user", {"columna__in": ["POR_HACER", "EN_PROCESO"]}),
                (AccionCAPA, "responsable_user", {"estado__in": ["PENDIENTE", "EN_PROCESO"]}),
                (Queja, "responsable_user", {"estado__in": ["ABIERTA", "EN_PROCESO"]}),
                (IndicadorKPI, "responsable_user", {}),
                (EvaluacionProveedor, "responsable_user", {"estado__in": ["EVALUADO", "SEGUIMIENTO"]}),
                (SalidaNoConforme, "responsable_user", {"estado__in": ["ABIERTA", "EN_TRATAMIENTO"]}),
            ]
            conteo: dict[int, int] = {}
            for modelo, campo, filtro in fuentes:
                base = (modelo.objects.filter(no_conformidad__empresa_id=emp_id)
                        if modelo is AccionCAPA else modelo.objects.filter(empresa_id=emp_id))
                qs = base.filter(**filtro).exclude(**{f"{campo}__isnull": True})
                for uid in qs.values_list(f"{campo}_id", flat=True):
                    if uid:
                        conteo[uid] = conteo.get(uid, 0) + 1
            carga = MiembroSerializer(miembros_empresa(empresa), many=True,
                                      context={"empresa_id": emp_id}).data
            for m in carga:
                m["asignados_abiertos"] = conteo.get(m["id"], 0)
            carga.sort(key=lambda x: x["asignados_abiertos"], reverse=True)

        return Response({
            "madurez": madurez, "nivel": nivel,
            "cumplimiento_iso": cumplimiento,
            "requisitos_evaluados": evaluados,
            "requisitos_total": RequisitoISO.objects.count(),
            "capitulos": capitulos,
            "dimensiones": dimensiones,
            "carga_equipo": carga,
        })

    @action(detail=False, methods=["get"])
    def tendencias(self, request):
        """Evolución histórica de la madurez del SGC + proyección a la meta +
        alertas inteligentes. Registra/actualiza el snapshot del día."""
        from datetime import timedelta
        from .models import SnapshotMadurez
        ids = self._emp(request)
        emp_id = ids[0] if ids else None
        if not emp_id:
            return Response({"detail": "Empresa inválida."}, status=400)
        hoy = date.today()

        # Recalcula los indicadores actuales (reusa lógica de panorama).
        pan = self.panorama(request).data
        nc = NoConformidad.objects.filter(empresa_id=emp_id)
        riesgos = list(Riesgo.objects.filter(empresa_id=emp_id))
        kpis = list(IndicadorKPI.objects.filter(empresa_id=emp_id))
        objs = ObjetivoCalidad.objects.filter(empresa_id=emp_id)
        quejas = Queja.objects.filter(empresa_id=emp_id).exclude(satisfaccion__isnull=True)
        sat = quejas.aggregate(p=Avg("satisfaccion"))["p"]

        # Registra/actualiza snapshot del día.
        SnapshotMadurez.objects.update_or_create(
            empresa_id=emp_id, fecha=hoy,
            defaults={
                "madurez": pan["madurez"], "cumplimiento_iso": pan["cumplimiento_iso"],
                "nc_abiertas": nc.filter(estado__in=["ABIERTA", "EN_PROCESO"]).count(),
                "nc_cerradas": nc.filter(estado="CERRADA").count(),
                "riesgos_altos": sum(1 for r in riesgos if r.severidad in ("ALTO", "CRITICO")),
                "kpis_en_meta": sum(1 for k in kpis if k.cumple), "kpis_total": len(kpis),
                "objetivos_logrados": objs.filter(estado="LOGRADO").count(),
                "satisfaccion": round(sat, 2) if sat else None,
            })

        # Serie histórica (últimos 12 snapshots).
        snaps = list(SnapshotMadurez.objects.filter(empresa_id=emp_id).order_by("fecha"))[-12:]
        serie = [{"fecha": str(s.fecha), "madurez": float(s.madurez),
                  "cumplimiento": float(s.cumplimiento_iso),
                  "nc_abiertas": s.nc_abiertas, "kpis_pct": round(s.kpis_en_meta / s.kpis_total * 100, 1) if s.kpis_total else 0}
                 for s in snaps]

        # Proyección lineal simple a 85% (meta de certificación).
        proyeccion = None
        if len(serie) >= 2:
            primero, ultimo = serie[0], serie[-1]
            d0 = date.fromisoformat(primero["fecha"])
            d1 = date.fromisoformat(ultimo["fecha"])
            dias = (d1 - d0).days or 1
            delta = ultimo["madurez"] - primero["madurez"]
            ritmo = delta / dias  # puntos por día
            actual = ultimo["madurez"]
            if ritmo > 0.001 and actual < 85:
                dias_faltan = (85 - actual) / ritmo
                proyeccion = {
                    "ritmo_mensual": round(ritmo * 30, 1),
                    "dias_estimados": int(dias_faltan),
                    "fecha_estimada": str(hoy + timedelta(days=int(dias_faltan))),
                    "meta": 85,
                }
            elif actual >= 85:
                proyeccion = {"listo": True, "meta": 85}
            else:
                proyeccion = {"ritmo_mensual": round(ritmo * 30, 1), "meta": 85}

        # Alertas inteligentes.
        alertas = []
        venc_pronto = nc.filter(estado__in=["ABIERTA", "EN_PROCESO"],
                                fecha_compromiso__gte=hoy,
                                fecha_compromiso__lte=hoy + timedelta(days=7))
        if venc_pronto.exists():
            alertas.append({"nivel": "alta", "tipo": "NC",
                            "mensaje": f"{venc_pronto.count()} no conformidad(es) vencen esta semana",
                            "ruta": "/sgc/no-conformidades"})
        nc_venc = nc.filter(estado__in=["ABIERTA", "EN_PROCESO"], fecha_compromiso__lt=hoy)
        if nc_venc.exists():
            alertas.append({"nivel": "critica", "tipo": "NC",
                            "mensaje": f"{nc_venc.count()} no conformidad(es) ya están vencidas",
                            "ruta": "/sgc/no-conformidades"})
        rc = sum(1 for r in riesgos if r.severidad == "CRITICO")
        if rc:
            alertas.append({"nivel": "alta", "tipo": "RIESGO",
                            "mensaje": f"{rc} riesgo(s) en nivel crítico sin controlar",
                            "ruta": "/sgc/riesgos"})
        equipos_venc = Equipo.objects.filter(empresa_id=emp_id, requiere_calibracion=True,
                                             activo=True, fecha_proxima_calibracion__lt=hoy).count()
        if equipos_venc:
            alertas.append({"nivel": "media", "tipo": "CALIBRACION",
                            "mensaje": f"{equipos_venc} equipo(s) con calibración vencida",
                            "ruta": "/sgc/equipos"})
        kpis_bajo = sum(1 for k in kpis if not k.cumple)
        if kpis_bajo:
            alertas.append({"nivel": "media", "tipo": "KPI",
                            "mensaje": f"{kpis_bajo} indicador(es) fuera de meta",
                            "ruta": "/sgc/kpis"})

        # Comparativa vs snapshot anterior (usa los objetos snapshot completos).
        comparativa = None
        if len(snaps) >= 2:
            a0, a1 = snaps[-2], snaps[-1]
            comparativa = {
                "madurez_delta": round(float(a1.madurez) - float(a0.madurez), 1),
                "cumplimiento_iso_delta": round(float(a1.cumplimiento_iso) - float(a0.cumplimiento_iso), 1),
                "nc_delta": a1.nc_abiertas - a0.nc_abiertas,
                "riesgos_altos_delta": a1.riesgos_altos - a0.riesgos_altos,
                "kpis_en_meta_delta": a1.kpis_en_meta - a0.kpis_en_meta,
            }

        return Response({
            "serie": serie, "proyeccion": proyeccion, "alertas": alertas,
            "comparativa": comparativa, "madurez_actual": pan["madurez"],
        })

    @action(detail=False, methods=["get"], url_path="analisis-nc")
    def analisis_nc(self, request):
        """Analítica de no conformidades: Pareto de categorías de causa, origen
        (distribución), tendencia mensual y tiempo promedio de cierre."""
        from datetime import timedelta
        ids = self._emp(request)
        ncs = list(NoConformidad.objects.filter(empresa_id__in=ids))
        total = len(ncs)

        # Pareto por categoría de causa (6M).
        cat_label = dict(NoConformidad._meta.get_field("categoria_causa").choices)
        cat_count = {}
        for n in ncs:
            if n.categoria_causa:
                cat_count[n.categoria_causa] = cat_count.get(n.categoria_causa, 0) + 1
        pareto = sorted(
            [{"categoria": k, "label": cat_label.get(k, k), "n": v} for k, v in cat_count.items()],
            key=lambda x: x["n"], reverse=True)
        acum = 0
        suma_cat = sum(x["n"] for x in pareto) or 1
        for x in pareto:
            acum += x["n"]
            x["acumulado_pct"] = round(acum / suma_cat * 100, 1)

        # Origen (distribución).
        ori_label = dict(NoConformidad.ORIGENES)
        ori_count = {}
        for n in ncs:
            ori_count[n.origen] = ori_count.get(n.origen, 0) + 1
        origen = [{"origen": k, "label": ori_label.get(k, k), "n": v}
                  for k, v in sorted(ori_count.items(), key=lambda x: x[1], reverse=True)]

        # Tendencia últimos 6 meses (detectadas vs cerradas).
        hoy = date.today()
        meses = []
        for i in range(5, -1, -1):
            y = hoy.year
            m = hoy.month - i
            while m <= 0:
                m += 12
                y -= 1
            meses.append((y, m))
        MN = ["", "Ene", "Feb", "Mar", "Abr", "May", "Jun", "Jul", "Ago", "Sep", "Oct", "Nov", "Dic"]
        tendencia = []
        for (y, m) in meses:
            det = sum(1 for n in ncs if n.fecha_deteccion and n.fecha_deteccion.year == y and n.fecha_deteccion.month == m)
            cer = sum(1 for n in ncs if n.fecha_cierre and n.fecha_cierre.year == y and n.fecha_cierre.month == m)
            tendencia.append({"mes": f"{MN[m]} {str(y)[2:]}", "detectadas": det, "cerradas": cer})

        # Tiempo promedio de cierre (días).
        dias = [(n.fecha_cierre - n.fecha_deteccion).days for n in ncs
                if n.fecha_cierre and n.fecha_deteccion]
        tiempo_cierre = round(sum(dias) / len(dias), 1) if dias else None

        abiertas = sum(1 for n in ncs if n.estado in ("ABIERTA", "EN_PROCESO"))
        cerradas = sum(1 for n in ncs if n.estado == "CERRADA")
        vencidas = sum(1 for n in ncs if n.estado in ("ABIERTA", "EN_PROCESO")
                       and n.fecha_compromiso and n.fecha_compromiso < hoy)
        return Response({
            "total": total, "abiertas": abiertas, "cerradas": cerradas, "vencidas": vencidas,
            "tiempo_cierre_promedio": tiempo_cierre,
            "tasa_cierre": round(cerradas / total * 100, 1) if total else 0,
            "pareto": pareto, "origen": origen, "tendencia": tendencia,
        })

    @action(detail=False, methods=["get"])
    def checklist_auditoria(self, request):
        """Checklist de auditoría por cláusula ISO para preparar la certificación.

        Por cada requisito de la norma devuelve: cláusula, título, estado de
        cumplimiento (del diagnóstico), evidencia declarada, hallazgos de
        auditoría ligados y un veredicto. Agrupado por capítulo (4-10)."""
        ids = self._emp(request)
        norma_id = request.query_params.get("norma")

        reqs = RequisitoISO.objects.select_related("norma").all()
        if norma_id:
            reqs = reqs.filter(norma_id=norma_id)
        reqs = list(reqs)

        # Evaluaciones del diagnóstico indexadas por requisito.
        evals = {e.requisito_id: e for e in
                 EvaluacionRequisito.objects.filter(empresa_id__in=ids, requisito__in=reqs)}
        # Hallazgos de auditoría por requisito.
        hallazgos_por_req: dict[int, list] = {}
        for h in (Hallazgo.objects.filter(auditoria__empresa_id__in=ids, requisito__in=reqs)
                  .select_related("auditoria")):
            hallazgos_por_req.setdefault(h.requisito_id, []).append({
                "tipo": h.tipo, "tipo_display": h.get_tipo_display(),
                "descripcion": h.descripcion, "auditoria": h.auditoria.titulo,
            })

        CAP_NOMBRE = {
            "4": "Contexto de la organización", "5": "Liderazgo", "6": "Planificación",
            "7": "Apoyo", "8": "Operación", "9": "Evaluación del desempeño", "10": "Mejora",
        }

        def clave_orden(clausula):
            partes = []
            for p in (clausula or "0").split("."):
                try:
                    partes.append(int(p))
                except ValueError:
                    partes.append(0)
            return partes

        reqs.sort(key=lambda r: clave_orden(r.clausula))
        por_cap: dict[str, dict] = {}
        veredictos = {"NO": "No conforme", "PARCIAL": "Conforme con observación",
                      "SI": "Conforme", "NA": "No aplica"}
        for r in reqs:
            cap = (r.clausula or "0").split(".")[0]
            ev = evals.get(r.id)
            cumple = ev.cumple if ev else "PENDIENTE"
            grupo = por_cap.setdefault(cap, {
                "capitulo": cap, "nombre": CAP_NOMBRE.get(cap, f"Capítulo {cap}"),
                "items": [], "conforme": 0, "total": 0,
            })
            item = {
                "id": r.id, "clausula": r.clausula, "titulo": r.titulo,
                "descripcion": r.descripcion,
                "cumple": cumple,
                "veredicto": veredictos.get(cumple, "Pendiente de evaluar"),
                "evidencia": ev.evidencia if ev else "",
                "observaciones": ev.observaciones if ev else "",
                "hallazgos": hallazgos_por_req.get(r.id, []),
            }
            grupo["items"].append(item)
            grupo["total"] += 1
            if cumple == "SI":
                grupo["conforme"] += 1

        capitulos = []
        for cap in sorted(por_cap, key=lambda c: int(c) if c.isdigit() else 99):
            g = por_cap[cap]
            g["cumplimiento"] = round(g["conforme"] / g["total"] * 100, 1) if g["total"] else 0
            capitulos.append(g)

        total = sum(g["total"] for g in capitulos)
        conformes = sum(g["conforme"] for g in capitulos)
        no_conformes = sum(1 for r in reqs if (evals.get(r.id) and evals[r.id].cumple == "NO"))
        observaciones = sum(1 for r in reqs if (evals.get(r.id) and evals[r.id].cumple == "PARCIAL"))
        pendientes = sum(1 for r in reqs if r.id not in evals)
        return Response({
            "capitulos": capitulos,
            "resumen": {
                "total": total, "conformes": conformes, "no_conformes": no_conformes,
                "observaciones": observaciones, "pendientes": pendientes,
                "cumplimiento": round(conformes / total * 100, 1) if total else 0,
                "listo_certificacion": no_conformes == 0 and pendientes == 0,
            },
        })


# ═════════════════════════════════════════════════════════════════════════════
# COLABORACIÓN MULTI-USUARIO (Pilar 1)
# ═════════════════════════════════════════════════════════════════════════════
class MiembrosViewSet(viewsets.ViewSet):
    """Usuarios de la empresa activa, para selectores de asignación/@menciones."""
    permission_classes = [permissions.IsAuthenticated]

    def list(self, request):
        ids = list(_empresas(request.user))
        emp = request.query_params.get("empresa")
        emp_id = int(emp) if (emp and int(emp) in ids) else (ids[0] if ids else None)
        if not emp_id:
            return Response([])
        empresa = Empresa.objects.filter(id=emp_id).first()
        if not empresa:
            return Response([])
        data = MiembroSerializer(miembros_empresa(empresa), many=True,
                                 context={"empresa_id": emp_id}).data
        return Response(data)

    @action(detail=False, methods=["get"])
    def carga(self, request):
        """Carga de trabajo del equipo del SGC: por cada usuario con acceso al
        módulo, cuántos registros tiene asignados y abiertos. Permite repartir
        el trabajo entre todos los miembros de forma equitativa (9.1)."""
        ids = list(_empresas(request.user))
        emp = request.query_params.get("empresa")
        emp_id = int(emp) if (emp and int(emp) in ids) else (ids[0] if ids else None)
        if not emp_id:
            return Response([])
        empresa = Empresa.objects.filter(id=emp_id).first()
        if not empresa:
            return Response([])

        # (modelo, campo_asignacion, filtro_abierto) por cada tipo asignable.
        fuentes = [
            (NoConformidad, "responsable_user", {"estado__in": ["ABIERTA", "EN_PROCESO"]}),
            (Riesgo, "responsable_user", {"estado__in": ["IDENTIFICADO", "EN_TRATAMIENTO"]}),
            (Auditoria, "auditor_lider_user", {"estado__in": ["PROGRAMADA", "EN_CURSO"]}),
            (Capacitacion, "responsable_user", {"estado__in": ["PROGRAMADA"]}),
            (Equipo, "responsable_user", {"activo": True}),
            (ObjetivoCalidad, "responsable_user", {"estado": "EN_CURSO"}),
            (TareaImplementacion, "responsable_user", {"columna__in": ["POR_HACER", "EN_PROCESO"]}),
            (AccionCAPA, "responsable_user", {"estado__in": ["PENDIENTE", "EN_PROCESO"]}),
            (Queja, "responsable_user", {"estado__in": ["ABIERTA", "EN_PROCESO"]}),
            (IndicadorKPI, "responsable_user", {}),
            (EvaluacionProveedor, "responsable_user", {"estado__in": ["EVALUADO", "SEGUIMIENTO"]}),
            (SalidaNoConforme, "responsable_user", {"estado__in": ["ABIERTA", "EN_TRATAMIENTO"]}),
        ]
        # AccionCAPA no tiene FK empresa directa; se filtra por la NC padre.
        conteo: dict[int, dict] = {}
        for modelo, campo, filtro_abierto in fuentes:
            if modelo is AccionCAPA:
                base = modelo.objects.filter(no_conformidad__empresa_id=emp_id)
            else:
                base = modelo.objects.filter(empresa_id=emp_id)
            qs = base.filter(**filtro_abierto).exclude(**{f"{campo}__isnull": True})
            for uid in qs.values_list(f"{campo}_id", flat=True):
                if uid is None:
                    continue
                conteo.setdefault(uid, {"asignados": 0, "por_tipo": {}})
                conteo[uid]["asignados"] += 1
                clave = modelo._meta.model_name
                conteo[uid]["por_tipo"][clave] = conteo[uid]["por_tipo"].get(clave, 0) + 1

        miembros = MiembroSerializer(miembros_empresa(empresa), many=True,
                                     context={"empresa_id": emp_id}).data
        for m in miembros:
            c = conteo.get(m["id"], {"asignados": 0, "por_tipo": {}})
            m["asignados_abiertos"] = c["asignados"]
            m["por_tipo"] = c["por_tipo"]
        miembros.sort(key=lambda x: x["asignados_abiertos"], reverse=True)
        return Response(miembros)


class ComentarioViewSet(viewsets.ModelViewSet):
    """Hilo de comentarios de cualquier registro del SGC (genérico vía ContentType)."""
    permission_classes = [permissions.IsAuthenticated]
    serializer_class = ComentarioSGCSerializer

    def get_queryset(self):
        qs = (ComentarioSGC.objects.filter(empresa__in=_empresas(self.request.user))
              .select_related("autor").prefetch_related("menciones"))
        tipo = self.request.query_params.get("tipo")
        objeto = self.request.query_params.get("objeto")
        if tipo and objeto:
            ct = ct_para(tipo)
            qs = qs.filter(content_type=ct, object_id=objeto) if ct else qs.none()
        elif tipo or objeto:
            qs = qs.none()
        return qs

    def _resolver(self, request):
        tipo = request.data.get("tipo")
        objeto_id = request.data.get("objeto")
        ct = ct_para(tipo)
        if not ct or not objeto_id:
            return None, None
        obj = ct.model_class().objects.filter(pk=objeto_id).first()
        if not obj or getattr(obj, "empresa_id", None) not in set(_empresas(request.user)):
            return None, None
        return ct, obj

    def create(self, request, *args, **kwargs):
        texto = (request.data.get("texto") or "").strip()
        ct, obj = self._resolver(request)
        if not obj or not texto:
            return Response({"detail": "tipo, objeto y texto son requeridos."}, status=400)
        c = ComentarioSGC.objects.create(
            empresa=obj.empresa, content_type=ct, object_id=obj.pk,
            autor=request.user, texto=texto,
        )
        menciones = extraer_menciones(texto, obj.empresa)
        if menciones:
            c.menciones.set(menciones)
        registrar_actividad(obj, request.user, "COMENTO", "comentó")
        # Notifica a responsables del objeto (sin las menciones, que reciben otra).
        ids_mencion = {u.pk for u in menciones}
        interesados = []
        for f in ("responsable_user", "auditor_lider_user", "creado_por"):
            u = getattr(obj, f, None)
            if u and u.pk not in ids_mencion:
                interesados.append(u)
        notificar(interesados, obj.empresa, request.user, "COMENTARIO",
                  "Nuevo comentario", mensaje=texto[:140], obj=obj)
        notificar(menciones, obj.empresa, request.user, "MENCION",
                  "Te mencionaron en un comentario", mensaje=texto[:140], obj=obj)
        return Response(ComentarioSGCSerializer(c).data, status=201)

    def perform_update(self, serializer):
        serializer.save(editado=True)

    def get_object(self):
        obj = super().get_object()
        # Solo el autor (o staff/superuser) puede editar/borrar su comentario.
        if self.request.method in ("PUT", "PATCH", "DELETE"):
            u = self.request.user
            if obj.autor_id != u.id and not (u.is_staff or u.is_superuser):
                from rest_framework.exceptions import PermissionDenied
                raise PermissionDenied("Solo el autor puede modificar el comentario.")
        return obj


class ActividadViewSet(viewsets.ReadOnlyModelViewSet):
    """Bitácora/timeline de un registro del SGC (solo lectura)."""
    permission_classes = [permissions.IsAuthenticated]
    serializer_class = ActividadSGCSerializer

    def get_queryset(self):
        qs = (ActividadSGC.objects.filter(empresa__in=_empresas(self.request.user))
              .select_related("actor"))
        tipo = self.request.query_params.get("tipo")
        objeto = self.request.query_params.get("objeto")
        if tipo and objeto:
            ct = ct_para(tipo)
            qs = qs.filter(content_type=ct, object_id=objeto) if ct else qs.none()
        elif tipo or objeto:
            qs = qs.none()
        return qs[:200]


class NotificacionViewSet(viewsets.ReadOnlyModelViewSet):
    """Notificaciones del usuario autenticado (campana del SGC)."""
    permission_classes = [permissions.IsAuthenticated]
    serializer_class = NotificacionCalidadSerializer

    def get_queryset(self):
        qs = NotificacionCalidad.objects.filter(destinatario=self.request.user).select_related("actor")
        if self.request.query_params.get("no_leidas") in ("1", "true", "True"):
            qs = qs.filter(leida=False)
        return qs

    @action(detail=False, methods=["get"])
    def conteo(self, request):
        n = NotificacionCalidad.objects.filter(destinatario=request.user, leida=False).count()
        return Response({"no_leidas": n})

    @action(detail=True, methods=["post"])
    def leer(self, request, pk=None):
        marcar_leidas(NotificacionCalidad.objects.filter(pk=pk, destinatario=request.user))
        return Response({"ok": True})

    @action(detail=False, methods=["post"])
    def marcar_todas(self, request):
        n = marcar_leidas(NotificacionCalidad.objects.filter(destinatario=request.user))
        return Response({"marcadas": n})


class MisPendientesViewSet(viewsets.ViewSet):
    """Bandeja unificada 'Mis pendientes de calidad' del usuario autenticado."""
    permission_classes = [permissions.IsAuthenticated]

    @action(detail=False, methods=["get"])
    def resumen(self, request):
        user = request.user
        ids = list(_empresas(user))
        emp = request.query_params.get("empresa")
        if emp and int(emp) in ids:
            ids = [int(emp)]
        hoy = date.today()
        items = []

        def push(tipo, obj, titulo, estado, fecha_limite, ruta, progreso=None, sub=None):
            vencido = bool(fecha_limite and fecha_limite < hoy and estado not in
                           ("CERRADA", "RESUELTA", "LOGRADO", "CONTROLADO", "IMPLEMENTADO", "VERIFICADO"))
            items.append({
                "tipo": tipo, "id": obj.pk, "titulo": titulo, "estado": estado,
                "fecha_limite": fecha_limite, "vencido": vencido,
                "url": f"{ruta}?focus={obj.pk}",
                "progreso": progreso, "sub": sub,
            })

        for nc in NoConformidad.objects.filter(empresa_id__in=ids, responsable_user=user,
                                               estado__in=["ABIERTA", "EN_PROCESO"]).prefetch_related("acciones"):
            acc = list(nc.acciones.all())
            sub = {"done": sum(1 for a in acc if a.estado in ("HECHA", "VERIFICADA")), "total": len(acc)} if acc else None
            push("No conformidad", nc, f"{nc.folio or 'NC'} · {nc.descripcion[:70]}",
                 nc.estado, nc.fecha_compromiso, "/sgc/no-conformidades", sub=sub)
        for r in Riesgo.objects.filter(empresa_id__in=ids, responsable_user=user,
                                       estado__in=["IDENTIFICADO", "EN_TRATAMIENTO"]):
            push("Riesgo", r, r.descripcion[:70], r.estado, None, "/sgc/riesgos")
        for o in ObjetivoCalidad.objects.filter(empresa_id__in=ids, responsable_user=user,
                                                estado="EN_CURSO").prefetch_related("tareas"):
            tareas = list(o.tareas.all())
            sub = {"done": sum(1 for t in tareas if t.avance >= 100), "total": len(tareas)} if tareas else None
            push("Objetivo", o, o.objetivo[:70], o.estado, o.fecha_limite, "/sgc/politica",
                 progreso=o.avance, sub=sub)
        for t in (TareaImplementacion.objects.filter(empresa_id__in=ids, responsable_user=user)
                  .exclude(columna="VERIFICADO").prefetch_related("subtareas")):
            subs = list(t.subtareas.all())
            sub = {"done": sum(1 for s in subs if s.avance >= 100), "total": len(subs)} if subs else None
            push("Implementación", t, t.titulo, t.get_columna_display(), t.fecha_limite,
                 "/sgc/implementacion", progreso=t.avance, sub=sub)
        for a in Auditoria.objects.filter(empresa_id__in=ids, auditor_lider_user=user,
                                          estado__in=["PROGRAMADA", "EN_CURSO"]):
            push("Auditoría", a, a.titulo, a.estado, a.fecha_programada, "/sgc/auditorias")
        for c in Capacitacion.objects.filter(empresa_id__in=ids, responsable_user=user,
                                             estado__in=["PROGRAMADA", "VENCIDA"]):
            push("Capacitación", c, c.curso, c.estado, c.fecha, "/sgc/capacitacion")
        for e in Equipo.objects.filter(empresa_id__in=ids, responsable_user=user,
                                       requiere_calibracion=True, activo=True):
            push("Calibración", e, f"{e.codigo} · {e.nombre}",
                 "Calibrar", e.fecha_proxima_calibracion, "/sgc/equipos")

        items.sort(key=lambda x: (not x["vencido"], x["fecha_limite"] or date.max))
        return Response({
            "total": len(items),
            "vencidos": sum(1 for i in items if i["vencido"]),
            "items": items,
        })


# ═════════════════════════════════════════════════════════════════════════════
# PILAR 2 · TABLERO DE IMPLEMENTACIÓN (Kanban ISO)
# ═════════════════════════════════════════════════════════════════════════════
class TareaImplementacionViewSet(_Colaborativo):
    queryset = TareaImplementacion.objects.select_related("responsable_user", "requisito")
    serializer_class = TareaImplementacionSerializer
    filterset_fields = ["empresa", "columna", "prioridad", "responsable_user", "requisito"]

    AVANCE_COL = {"POR_HACER": 0, "EN_PROCESO": 50, "IMPLEMENTADO": 90, "VERIFICADO": 100}

    @action(detail=False, methods=["get"])
    def tablero(self, request):
        ids = list(_empresas(request.user))
        emp = request.query_params.get("empresa")
        if emp and int(emp) in ids:
            ids = [int(emp)]
        tareas = self.get_queryset().filter(empresa_id__in=ids)
        cols = {c: [] for c, _ in TareaImplementacion.COLUMNAS}
        for t in tareas:
            cols.setdefault(t.columna, []).append(TareaImplementacionSerializer(t).data)
        total = tareas.count()
        verificadas = tareas.filter(columna="VERIFICADO").count()
        return Response({
            "columnas": [{"clave": c, "titulo": lbl, "tareas": cols.get(c, [])}
                         for c, lbl in TareaImplementacion.COLUMNAS],
            "total": total, "verificadas": verificadas,
            "progreso": round(verificadas / total * 100, 1) if total else 0,
        })

    @action(detail=True, methods=["post"])
    def mover(self, request, pk=None):
        t = self.get_object()
        antes = t.columna
        col = request.data.get("columna")
        if col not in dict(TareaImplementacion.COLUMNAS):
            return Response({"detail": "Columna inválida."}, status=400)
        t.columna = col
        if "orden" in request.data:
            t.orden = request.data["orden"]
        # El avance sigue a la columna salvo que se especifique a mano.
        t.avance = request.data.get("avance", self.AVANCE_COL.get(col, t.avance))
        t.save()
        if antes != col:
            registrar_actividad(t, request.user, "CAMBIO_ESTADO",
                                f"movió la tarjeta a {t.get_columna_display()}",
                                {"columna": [antes, col]})
            if t.responsable_user:
                notificar([t.responsable_user], t.empresa, request.user, "CAMBIO_ESTADO",
                          f"{t.titulo}: {antes} → {col}", obj=t)
        return Response(TareaImplementacionSerializer(t).data)

    @action(detail=False, methods=["post"])
    def generar(self, request):
        """Crea una tarjeta por cada cláusula ISO que aún no tenga tarea."""
        ids = list(_empresas(request.user))
        emp_id = request.data.get("empresa") or (ids[0] if ids else None)
        if not emp_id or int(emp_id) not in ids:
            return Response({"detail": "Empresa inválida."}, status=400)
        existentes = set(TareaImplementacion.objects.filter(empresa_id=emp_id)
                         .exclude(requisito=None).values_list("requisito_id", flat=True))
        creadas = 0
        for req in RequisitoISO.objects.select_related("norma"):
            if req.id in existentes:
                continue
            TareaImplementacion.objects.create(
                empresa_id=emp_id, requisito=req,
                titulo=f"{req.clausula} · {req.titulo}",
                descripcion=f"Implementar el requisito {req.clausula} de {req.norma.codigo}.",
                columna="POR_HACER", creado_por=request.user,
            )
            creadas += 1
        return Response({"creadas": creadas})


# ═════════════════════════════════════════════════════════════════════════════
# PILAR 3 · CAPA AVANZADO (acciones múltiples + evidencias)
# ═════════════════════════════════════════════════════════════════════════════
class AccionCAPAViewSet(viewsets.ModelViewSet):
    permission_classes = [permissions.IsAuthenticated]
    serializer_class = AccionCAPASerializer
    filterset_fields = ["no_conformidad", "estado", "responsable_user"]

    def get_queryset(self):
        return (AccionCAPA.objects.filter(no_conformidad__empresa__in=_empresas(self.request.user))
                .select_related("responsable_user", "no_conformidad"))

    @staticmethod
    def _avanzar_nc(nc):
        """Auto-progreso: con acciones definidas, una NC ABIERTA pasa a EN_PROCESO."""
        if nc.estado == "ABIERTA" and nc.acciones.exists():
            nc.estado = "EN_PROCESO"
            nc.save(update_fields=["estado"])
            return True
        return False

    def perform_create(self, serializer):
        accion = serializer.save()
        nc = accion.no_conformidad
        registrar_actividad(nc, self.request.user, "ACTUALIZO",
                            f"agregó una acción: {accion.descripcion[:60]}")
        self._avanzar_nc(nc)
        if accion.responsable_user:
            notificar([accion.responsable_user], nc.empresa, self.request.user, "ASIGNACION",
                      f"Acción CAPA asignada en {nc.folio or nc.id}", mensaje=accion.descripcion[:140], obj=nc)

    def perform_update(self, serializer):
        antes = serializer.instance.estado
        accion = serializer.save()
        if accion.estado != antes:
            registrar_actividad(accion.no_conformidad, self.request.user, "CAMBIO_ESTADO",
                                f"acción «{accion.descripcion[:40]}» → {accion.get_estado_display()}")
        self._avanzar_nc(accion.no_conformidad)


class EvidenciaViewSet(viewsets.ModelViewSet):
    """Adjuntos de evidencia (genéricos) para cualquier registro del SGC."""
    permission_classes = [permissions.IsAuthenticated]
    serializer_class = EvidenciaSGCSerializer

    def get_queryset(self):
        qs = EvidenciaSGC.objects.filter(empresa__in=_empresas(self.request.user)).select_related("subido_por")
        tipo = self.request.query_params.get("tipo")
        objeto = self.request.query_params.get("objeto")
        if tipo and objeto:
            ct = ct_para(tipo)
            qs = qs.filter(content_type=ct, object_id=objeto) if ct else qs.none()
        elif tipo or objeto:
            qs = qs.none()
        return qs

    def create(self, request, *args, **kwargs):
        tipo = request.data.get("tipo")
        objeto_id = request.data.get("objeto")
        ct = ct_para(tipo)
        if not ct or not objeto_id or "archivo" not in request.FILES:
            return Response({"detail": "tipo, objeto y archivo son requeridos."}, status=400)
        obj = ct.model_class().objects.filter(pk=objeto_id).first()
        if not obj or getattr(obj, "empresa_id", None) not in set(_empresas(request.user)):
            return Response({"detail": "Objeto no encontrado."}, status=404)
        archivo = request.FILES["archivo"]
        ev = EvidenciaSGC.objects.create(
            empresa=obj.empresa, content_type=ct, object_id=obj.pk, archivo=archivo,
            nombre=request.data.get("nombre") or archivo.name,
            descripcion=request.data.get("descripcion", ""), subido_por=request.user,
        )
        registrar_actividad(obj, request.user, "ADJUNTO", f"adjuntó evidencia: {ev.nombre}")
        return Response(EvidenciaSGCSerializer(ev, context={"request": request}).data, status=201)


# ═════════════════════════════════════════════════════════════════════════════
# PILAR 4 · PROCESOS (SIPOC) + MEDICIONES KPI
# ═════════════════════════════════════════════════════════════════════════════
class ProcesoViewSet(_Colaborativo):
    asignacion_field = "dueno_user"
    queryset = Proceso.objects.select_related("dueno_user").prefetch_related("kpis", "riesgos", "objetivos")
    serializer_class = ProcesoSerializer
    filterset_fields = ["empresa", "tipo", "activo", "dueno_user"]

    @action(detail=True, methods=["get"])
    def panorama(self, request, pk=None):
        """Tablero 360° del proceso: KPIs, riesgos y objetivos vinculados."""
        from .models import IndicadorKPI as _KPI, ObjetivoCalidad as _Obj, Riesgo as _Rie
        p = self.get_object()
        kpis = list(p.kpis.all())
        riesgos = list(p.riesgos.all())
        objetivos = list(p.objetivos.all())
        return Response({
            "proceso": ProcesoSerializer(p).data,
            "kpis": IndicadorKPISerializer(kpis, many=True).data,
            "riesgos": RiesgoSerializer(riesgos, many=True).data,
            "objetivos": ObjetivoCalidadSerializer(objetivos, many=True).data,
            "resumen": {
                "kpis_total": len(kpis),
                "kpis_en_meta": sum(1 for k in kpis if k.cumple),
                "riesgos_total": len(riesgos),
                "riesgos_altos": sum(1 for r in riesgos if r.severidad in ("ALTO", "CRITICO")),
                "objetivos_total": len(objetivos),
                "objetivos_logrados": sum(1 for o in objetivos if o.estado == "LOGRADO"),
            },
        })


class MedicionKPIViewSet(viewsets.ModelViewSet):
    permission_classes = [permissions.IsAuthenticated]
    serializer_class = MedicionKPISerializer
    filterset_fields = ["kpi"]

    def get_queryset(self):
        return MedicionKPI.objects.filter(kpi__empresa__in=_empresas(self.request.user)).select_related("kpi")

    def perform_create(self, serializer):
        medicion = serializer.save(registrado_por=self.request.user)
        # Actualiza el valor actual del KPI con la última medición.
        kpi = medicion.kpi
        ultima = kpi.mediciones.order_by("-fecha").first()
        if ultima and ultima.id == medicion.id:
            kpi.valor_actual = medicion.valor
            kpi.save(update_fields=["valor_actual"])


class TareaObjetivoViewSet(viewsets.ModelViewSet):
    """Tareas/resultados clave de un objetivo. Cada cambio recalcula el avance
    del objetivo padre y lo autocompleta al llegar al 100 %."""
    permission_classes = [permissions.IsAuthenticated]
    serializer_class = TareaObjetivoSerializer
    filterset_fields = ["objetivo"]

    def get_queryset(self):
        return (TareaObjetivo.objects.filter(objetivo__empresa__in=_empresas(self.request.user))
                .select_related("responsable_user", "objetivo"))

    def _recomputar(self, objetivo):
        objetivo.recomputar()
        return {"objetivo_avance": objetivo.avance, "objetivo_estado": objetivo.estado}

    def create(self, request, *args, **kwargs):
        resp = super().create(request, *args, **kwargs)
        obj = self._obj
        extra = self._recomputar(obj)
        resp.data.update(extra)
        if obj.responsable_user and obj.responsable_user_id != request.user.id:
            notificar([obj.responsable_user], obj.empresa, request.user, "INFO",
                      f"Nueva tarea en objetivo: {obj.objetivo[:60]}", obj=obj)
        return resp

    def perform_create(self, serializer):
        tarea = serializer.save()
        self._obj = tarea.objetivo

    def update(self, request, *args, **kwargs):
        resp = super().update(request, *args, **kwargs)
        obj = self.get_object().objetivo
        resp.data.update(self._recomputar(obj))
        return resp

    def destroy(self, request, *args, **kwargs):
        tarea = self.get_object()
        obj = tarea.objetivo
        tarea.delete()
        obj.recomputar()
        return Response({"objetivo_avance": obj.avance, "objetivo_estado": obj.estado})


class SubtareaImplementacionViewSet(viewsets.ModelViewSet):
    """Subtareas de una tarjeta del tablero; recalculan el avance/columna del padre."""
    permission_classes = [permissions.IsAuthenticated]
    serializer_class = SubtareaImplementacionSerializer
    filterset_fields = ["tarea"]

    def get_queryset(self):
        return (SubtareaImplementacion.objects.filter(tarea__empresa__in=_empresas(self.request.user))
                .select_related("responsable_user", "tarea"))

    @staticmethod
    def _extra(t):
        return {"tarea_avance": t.avance, "tarea_columna": t.columna}

    def create(self, request, *args, **kwargs):
        resp = super().create(request, *args, **kwargs)
        self._t.recomputar()
        resp.data.update(self._extra(self._t))
        return resp

    def perform_create(self, serializer):
        sub = serializer.save()
        self._t = sub.tarea

    def update(self, request, *args, **kwargs):
        resp = super().update(request, *args, **kwargs)
        t = self.get_object().tarea
        t.recomputar()
        resp.data.update(self._extra(t))
        return resp

    def destroy(self, request, *args, **kwargs):
        sub = self.get_object()
        t = sub.tarea
        sub.delete()
        t.recomputar()
        return Response(self._extra(t))


class AcuerdoRevisionViewSet(viewsets.ModelViewSet):
    """Acuerdos (salidas 9.3.3) de la revisión, asignables y con seguimiento."""
    permission_classes = [permissions.IsAuthenticated]
    serializer_class = AcuerdoRevisionSerializer
    filterset_fields = ["revision", "estado"]

    def get_queryset(self):
        return (AcuerdoRevision.objects.filter(revision__empresa__in=_empresas(self.request.user))
                .select_related("responsable_user", "revision"))

    def perform_create(self, serializer):
        ac = serializer.save()
        if ac.responsable_user and ac.responsable_user_id != self.request.user.id:
            notificar([ac.responsable_user], ac.revision.empresa, self.request.user, "ASIGNACION",
                      f"Acuerdo de revisión: {ac.descripcion[:60]}", obj=ac.revision)


# ═════════════════════════════════════════════════════════════════════════════
# ENCUESTAS PÚBLICAS
# ═════════════════════════════════════════════════════════════════════════════
class EncuestaViewSet(_Scoped):
    queryset = Encuesta.objects.all()
    serializer_class = EncuestaSerializer
    filterset_fields = ["empresa", "activa"]

    def perform_create(self, serializer):
        serializer.save(token=secrets.token_urlsafe(9), creado_por=self.request.user)

    @action(detail=True, methods=["get"])
    def resultados(self, request, pk=None):
        enc = self.get_object()
        resp = list(enc.respuestas.all())
        sats = [r.satisfaccion for r in resp if r.satisfaccion]
        npss = [r.nps for r in resp if r.nps is not None]
        prom = round(sum(sats) / len(sats), 1) if sats else None
        csat = round(sum(1 for s in sats if s >= 4) / len(sats) * 100) if sats else None
        # NPS = %promotores(9-10) - %detractores(0-6)
        nps_val = None
        if npss:
            promo = sum(1 for n in npss if n >= 9); detr = sum(1 for n in npss if n <= 6)
            nps_val = round((promo - detr) / len(npss) * 100)
        dist = [{"estrella": s, "n": sum(1 for x in sats if x == s)} for s in range(1, 6)]
        return Response({
            "total": len(resp), "satisfaccion_prom": prom, "csat": csat, "nps": nps_val,
            "distribucion": dist,
            "respuestas": RespuestaEncuestaSerializer(resp, many=True).data,
        })


@api_view(["GET"])
@permission_classes([AllowAny])
def encuesta_publica(request, token):
    enc = Encuesta.objects.filter(token=token).first()
    if not enc:
        return Response({"detail": "Encuesta no encontrada."}, status=404)
    if not enc.activa:
        return Response({"detail": "Esta encuesta está cerrada.", "cerrada": True}, status=200)
    return Response(EncuestaPublicaSerializer(enc).data)


@api_view(["POST"])
@permission_classes([AllowAny])
def responder_encuesta(request, token):
    enc = Encuesta.objects.filter(token=token, activa=True).first()
    if not enc:
        return Response({"detail": "Encuesta no disponible."}, status=404)
    data = request.data or {}
    respuestas = data.get("respuestas") or {}
    # Deriva satisfacción (primera pregunta rating) y NPS (primera nps).
    sat = nps = None
    for p in enc.preguntas:
        pid = str(p.get("id"))
        val = respuestas.get(pid)
        if val in (None, ""):
            continue
        if p.get("tipo") == "rating" and sat is None:
            try: sat = int(val)
            except (TypeError, ValueError): pass
        if p.get("tipo") == "nps" and nps is None:
            try: nps = int(val)
            except (TypeError, ValueError): pass
    r = RespuestaEncuesta.objects.create(
        encuesta=enc, respuestas=respuestas,
        cliente_nombre=data.get("cliente_nombre", "")[:200],
        cliente_email=data.get("cliente_email", "")[:200],
        satisfaccion=sat, nps=nps,
    )
    # Crea un caso en Quejas con cada respuesta (voz del cliente) si está activado.
    if enc.crear_queja:
        textos = []
        for p in enc.preguntas:
            v = respuestas.get(str(p.get("id")))
            if v not in (None, "", []):
                textos.append(f"{p.get('titulo')}: {v if not isinstance(v, list) else ', '.join(map(str, v))}")
        tipo = "FELICITACION" if (sat and sat >= 4) else "SUGERENCIA" if (sat == 3) else "QUEJA"
        if sat is None and nps is not None:
            tipo = "FELICITACION" if nps >= 9 else "QUEJA" if nps <= 6 else "SUGERENCIA"
        q = Queja.objects.create(
            empresa=enc.empresa, tipo=tipo,
            cliente=data.get("cliente_nombre", "") or "Anónimo",
            descripcion=f"[Encuesta: {enc.titulo}]\n" + "\n".join(textos[:20]),
            satisfaccion=sat, estado="ABIERTA",
        )
        r.queja = q; r.save(update_fields=["queja"])
        # Notifica a calidad/dirección si fue una queja/insatisfacción.
        if tipo == "QUEJA":
            from apps.core.models import UsuarioEmpresa
            uids = UsuarioEmpresa.objects.filter(empresa=enc.empresa, activo=True, rol__in=["OWNER", "STAFF", "MANAGER"]).values_list("user_id", flat=True)
            from django.contrib.auth import get_user_model
            notificar(list(get_user_model().objects.filter(id__in=uids)), enc.empresa, None, "VENCIMIENTO",
                      f"Nueva queja por encuesta: {enc.titulo}", mensaje=(q.cliente or "Cliente"), url="/sgc/quejas")
    return Response({"ok": True, "mensaje": enc.mensaje_gracias})


# ═════════════════════════════════════════════════════════════════════════════
# COMPETENCIAS GRANULARES (7.2)
# ═════════════════════════════════════════════════════════════════════════════
class CompetenciaPerfilViewSet(viewsets.ModelViewSet):
    permission_classes = [permissions.IsAuthenticated]
    serializer_class = CompetenciaPerfilSerializer
    filterset_fields = ["perfil", "tipo"]

    def get_queryset(self):
        return CompetenciaPerfil.objects.filter(perfil__empresa__in=_empresas(self.request.user)).select_related("perfil")


class EvaluacionDetalleViewSet(viewsets.ModelViewSet):
    permission_classes = [permissions.IsAuthenticated]
    serializer_class = EvaluacionDetalleSerializer
    filterset_fields = ["evaluacion"]

    def get_queryset(self):
        return EvaluacionDetalle.objects.filter(evaluacion__empresa__in=_empresas(self.request.user)).select_related("evaluacion")

    def perform_create(self, serializer):
        d = serializer.save(); d.evaluacion.recomputar()

    def perform_update(self, serializer):
        d = serializer.save(); d.evaluacion.recomputar()

    def perform_destroy(self, instance):
        ev = instance.evaluacion; instance.delete(); ev.recomputar()


# ═════════════════════════════════════════════════════════════════════════════
# ECOSISTEMA DE CERTIFICACIÓN
# ═════════════════════════════════════════════════════════════════════════════

# Roadmap estándar de certificación ISO 9001 (fase, título, cláusula, ruta).
ROADMAP_BASE = [
    ("DIAGNOSTICO", "Realizar el diagnóstico inicial (gap analysis)", "4-10", "/sgc/diagnostico"),
    ("DIAGNOSTICO", "Definir el contexto y partes interesadas (FODA)", "4.1", "/sgc/contexto"),
    ("DIAGNOSTICO", "Determinar el alcance del SGC", "4.3", "/sgc/contexto"),
    ("PLANEACION", "Redactar la política de calidad", "5.2", "/sgc/politica"),
    ("PLANEACION", "Establecer objetivos de calidad medibles", "6.2", "/sgc/politica"),
    ("PLANEACION", "Identificar riesgos y oportunidades", "6.1", "/sgc/riesgos"),
    ("PLANEACION", "Mapear los procesos (SIPOC)", "4.4", "/sgc/procesos"),
    ("DOCUMENTACION", "Generar la documentación con plantillas", "7.5", "/sgc/plantillas"),
    ("DOCUMENTACION", "Definir la lista maestra de registros", "7.5.3", "/sgc/registros"),
    ("DOCUMENTACION", "Establecer la matriz de comunicación", "7.4", "/sgc/comunicacion"),
    ("IMPLEMENTACION", "Capacitar al personal y evaluar competencias", "7.2", "/sgc/competencias"),
    ("IMPLEMENTACION", "Avanzar el tablero de implementación", "8.1", "/sgc/implementacion"),
    ("IMPLEMENTACION", "Controlar equipos y calibraciones", "7.1.5", "/sgc/equipos"),
    ("MEDICION", "Definir y medir indicadores KPI", "9.1", "/sgc/kpis"),
    ("MEDICION", "Medir la satisfacción del cliente", "9.1.2", "/sgc/quejas"),
    ("MEDICION", "Evaluar a los proveedores", "8.4", "/sgc/evaluacion-proveedores"),
    ("AUDITORIA_INTERNA", "Crear el programa anual de auditorías", "9.2", "/sgc/programa-auditorias"),
    ("AUDITORIA_INTERNA", "Ejecutar la auditoría interna", "9.2", "/sgc/auditorias"),
    ("AUDITORIA_INTERNA", "Levantar y cerrar no conformidades", "10.2", "/sgc/no-conformidades"),
    ("REVISION_DIRECCION", "Realizar la revisión por la dirección", "9.3", "/sgc/revision-direccion"),
    ("PREAUDITORIA", "Verificar el checklist de auditoría (modo auditor)", "9.2", "/sgc/auditor"),
    ("PREAUDITORIA", "Cerrar acciones correctivas pendientes", "10.2", "/sgc/no-conformidades"),
    ("CERTIFICACION", "Seleccionar organismo certificador", "-", ""),
    ("CERTIFICACION", "Aprobar la auditoría de certificación", "-", ""),
]


# ── Generación de Word profesional (docx) ────────────────────────────────────
_DOCX_INDIGO = (0x4F, 0x46, 0xE5)
_DOCX_DARK = (0x1E, 0x29, 0x3B)
_DOCX_GREY = (0x6B, 0x72, 0x80)
_DOCX_FILL_HEAD = "4F46E5"   # cabecera fuerte (texto blanco)
_DOCX_FILL_SOFT = "EEF0FB"   # relleno suave (cabeceras de tabla del cuerpo)


def _docx_shade(cell, hexfill):
    from docx.oxml.ns import qn
    from docx.oxml import OxmlElement
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hexfill)
    tcPr.append(shd)


def _docx_set_cell(cell, text, *, bold=False, size=9, color=None, align=None, fill=None):
    from docx.shared import Pt, RGBColor
    from docx.enum.table import WD_ALIGN_VERTICAL
    cell.text = ""
    p = cell.paragraphs[0]
    if align is not None:
        p.alignment = align
    r = p.add_run("" if text is None else str(text))
    r.bold = bold
    r.font.size = Pt(size)
    if color is not None:
        r.font.color.rgb = RGBColor(*color)
    if fill:
        _docx_shade(cell, fill)
    try:
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
    except Exception:
        pass
    return r


def _docx_para_rule(p, color="4F46E5", sz=6):
    """Línea inferior bajo un encabezado (regla de color)."""
    from docx.oxml.ns import qn
    from docx.oxml import OxmlElement
    pPr = p._p.get_or_add_pPr()
    pbdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), str(sz))
    bottom.set(qn("w:space"), "3")
    bottom.set(qn("w:color"), color)
    pbdr.append(bottom)
    pPr.append(pbdr)


def _docx_field(par, instr):
    """Inserta un campo de Word (p. ej. PAGE / NUMPAGES) en un párrafo."""
    from docx.oxml.ns import qn
    from docx.oxml import OxmlElement
    run = par.add_run()
    b = OxmlElement("w:fldChar"); b.set(qn("w:fldCharType"), "begin")
    it = OxmlElement("w:instrText"); it.set(qn("xml:space"), "preserve"); it.text = instr
    e = OxmlElement("w:fldChar"); e.set(qn("w:fldCharType"), "end")
    run._r.append(b); run._r.append(it); run._r.append(e)


def _docx_render_body(doc, texto):
    """Renderiza el cuerpo (texto con secciones y tablas ASCII) como contenido
    Word profesional: encabezados con regla, subtítulos, viñetas y tablas con
    cabecera sombreada y bordes."""
    import re as _re
    from docx.shared import Pt, RGBColor
    from docx.enum.text import WD_ALIGN_PARAGRAPH

    INDIGO = RGBColor(*_DOCX_INDIGO)
    DARK = RGBColor(*_DOCX_DARK)
    lines = texto.split("\n")
    tbuf: list = []

    def flush():
        rows = [r for r in tbuf if not _re.match(r"^[\s\-|]+$", r.strip())]
        tbuf.clear()
        if not rows:
            return
        celdas = [[c.strip() for c in r.strip().strip("|").split("|")] for r in rows]
        ncol = max(len(c) for c in celdas)
        t = doc.add_table(rows=len(celdas), cols=ncol)
        t.style = "Table Grid"
        for i, fila in enumerate(celdas):
            for j in range(ncol):
                cell = t.rows[i].cells[j]
                if i == 0:
                    _docx_set_cell(cell, fila[j] if j < len(fila) else "", bold=True,
                                   size=8.5, color=_DOCX_INDIGO, fill=_DOCX_FILL_SOFT)
                else:
                    _docx_set_cell(cell, fila[j] if j < len(fila) else "", size=8.5)
        doc.add_paragraph("")

    for raw in lines:
        s = raw.strip()
        if "|" in s:
            tbuf.append(raw)
            continue
        if tbuf:
            flush()
        if not s:
            continue
        if _re.match(r"^[═─\-_]{3,}$", s):
            continue
        # Sección numerada principal "1. OBJETIVO" → encabezado con regla.
        if _re.match(r"^\d+\.\s+[A-ZÁÉÍÓÚÑ]", s):
            h = doc.add_paragraph()
            r = h.add_run(s); r.bold = True; r.font.size = Pt(12.5); r.font.color.rgb = INDIGO
            _docx_para_rule(h)
            continue
        # Subsección "4.1 ..." → subtítulo en negrita.
        if _re.match(r"^\d+\.\d+\S*\s+\S", s):
            h = doc.add_paragraph()
            r = h.add_run(s); r.bold = True; r.font.size = Pt(10.5); r.font.color.rgb = DARK
            continue
        # Título en mayúsculas (formatos/matrices).
        if _re.match(r"^[A-ZÁÉÍÓÚÑ][A-ZÁÉÍÓÚÑ0-9 ,/()\-]{3,}$", s) and not _re.search(r"[a-záéíóúñ]", s):
            h = doc.add_paragraph()
            r = h.add_run(s); r.bold = True; r.font.size = Pt(11); r.font.color.rgb = INDIGO
            continue
        # Viñetas y literales.
        if _re.match(r"^([•\-]|[a-z]\)|\d+\))\s+", s):
            p = doc.add_paragraph(s, style="List Bullet")
            continue
        doc.add_paragraph(s)
    if tbuf:
        flush()


def _docx_controlado(*, empresa, titulo, subtitulo, codigo, version, fecha,
                     clasificacion, body, elaboro=" ", reviso=" ", aprobo=" ",
                     historial=None):
    """Construye un documento Word controlado (BytesIO) de aspecto profesional:
    encabezado de control documental, control de cambios, firmas, cuerpo con
    tablas/encabezados nativos, y encabezado/pie con numeración de página."""
    from io import BytesIO
    from docx import Document
    from docx.shared import Pt, RGBColor, Inches
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.enum.table import WD_TABLE_ALIGNMENT

    WHITE = (0xFF, 0xFF, 0xFF)
    CENTER = WD_ALIGN_PARAGRAPH.CENTER
    LEFT = WD_ALIGN_PARAGRAPH.LEFT
    RIGHT = WD_ALIGN_PARAGRAPH.RIGHT

    doc = Document()
    sec = doc.sections[0]
    sec.left_margin = Inches(0.9); sec.right_margin = Inches(0.9)
    sec.top_margin = Inches(0.7); sec.bottom_margin = Inches(0.75)
    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(10.5)

    # Encabezado de página (código + título).
    hp = sec.header.paragraphs[0]; hp.text = ""; hp.alignment = RIGHT
    rr = hp.add_run(f"{codigo}   ·   {titulo}"); rr.font.size = Pt(7.5); rr.font.color.rgb = RGBColor(*_DOCX_GREY)

    # Pie de página (leyenda + numeración).
    fp = sec.footer.paragraphs[0]; fp.text = ""; fp.alignment = CENTER
    r = fp.add_run("Documento controlado conforme a ISO 9001:2015 (7.5). Una vez impreso es copia NO controlada.    ")
    r.font.size = Pt(7); r.font.color.rgb = RGBColor(*_DOCX_GREY)
    r = fp.add_run("Página "); r.font.size = Pt(7); r.font.color.rgb = RGBColor(*_DOCX_GREY)
    _docx_field(fp, "PAGE")
    r = fp.add_run(" de "); r.font.size = Pt(7); r.font.color.rgb = RGBColor(*_DOCX_GREY)
    _docx_field(fp, "NUMPAGES")

    # ── Bloque de encabezado de control documental ──
    ht = doc.add_table(rows=4, cols=4); ht.style = "Table Grid"; ht.alignment = WD_TABLE_ALIGNMENT.CENTER
    a = ht.cell(0, 0).merge(ht.cell(0, 1))
    _docx_set_cell(a, empresa, bold=True, size=12, color=_DOCX_DARK, align=LEFT)
    b = ht.cell(0, 2).merge(ht.cell(0, 3))
    _docx_set_cell(b, "Sistema de Gestión de Calidad · ISO 9001:2015", bold=True, size=8, color=_DOCX_INDIGO, align=RIGHT)
    tcell = ht.cell(1, 0).merge(ht.cell(1, 1)).merge(ht.cell(1, 2)).merge(ht.cell(1, 3))
    _docx_set_cell(tcell, (titulo or "").upper(), bold=True, size=16, color=_DOCX_INDIGO, align=CENTER)
    if subtitulo:
        sp = tcell.add_paragraph(); sp.alignment = CENTER
        sr = sp.add_run(subtitulo); sr.italic = True; sr.font.size = Pt(8.5); sr.font.color.rgb = RGBColor(*_DOCX_GREY)
    for j, l in enumerate(["Código", "Versión", "Fecha de emisión", "Clasificación"]):
        _docx_set_cell(ht.cell(2, j), l, bold=True, size=8, color=WHITE, align=CENTER, fill=_DOCX_FILL_HEAD)
    for j, v in enumerate([codigo or "—", version or "1.0", fecha, clasificacion]):
        _docx_set_cell(ht.cell(3, j), v, size=9, align=CENTER)
    doc.add_paragraph("")

    # ── Cuerpo del documento ──
    _docx_render_body(doc, body)

    # ── Control de cambios ──
    doc.add_paragraph("")
    h = doc.add_paragraph(); r = h.add_run("CONTROL DE CAMBIOS")
    r.bold = True; r.font.size = Pt(11); r.font.color.rgb = RGBColor(*_DOCX_INDIGO); _docx_para_rule(h)
    hist = historial or [(version or "1.0", fecha, "Emisión inicial", elaboro)]
    ct = doc.add_table(rows=1 + len(hist), cols=4); ct.style = "Table Grid"
    for j, hh in enumerate(["Versión", "Fecha", "Descripción del cambio", "Autor"]):
        _docx_set_cell(ct.cell(0, j), hh, bold=True, size=8, color=WHITE, align=CENTER, fill=_DOCX_FILL_HEAD)
    for i, fila in enumerate(hist, start=1):
        for j, val in enumerate(fila):
            _docx_set_cell(ct.cell(i, j), val, size=8.5)

    # ── Firmas de control ──
    doc.add_paragraph("")
    h = doc.add_paragraph(); r = h.add_run("APROBACIÓN")
    r.bold = True; r.font.size = Pt(11); r.font.color.rgb = RGBColor(*_DOCX_INDIGO); _docx_para_rule(h)
    ft = doc.add_table(rows=3, cols=3); ft.style = "Table Grid"
    for j, l in enumerate(["Elaboró", "Revisó", "Aprobó"]):
        _docx_set_cell(ft.cell(0, j), l, bold=True, size=8, color=WHITE, align=CENTER, fill=_DOCX_FILL_HEAD)
    for j, nm in enumerate([elaboro, reviso, aprobo]):
        _docx_set_cell(ft.cell(1, j), nm or " ", size=9, align=CENTER)
    for j in range(3):
        _docx_set_cell(ft.cell(2, j), "Nombre y firma", size=7.5, color=_DOCX_GREY, align=CENTER)

    bio = BytesIO(); doc.save(bio); bio.seek(0)
    return bio


class PlantillaDocumentoViewSet(viewsets.ModelViewSet):
    """Biblioteca de plantillas ISO. Incluye globales (empresa null) + propias."""
    permission_classes = [permissions.IsAuthenticated]
    serializer_class = PlantillaDocumentoSerializer
    filterset_fields = ["categoria", "clausula", "obligatoria"]

    def get_queryset(self):
        ids = list(_empresas(self.request.user))
        return PlantillaDocumento.objects.filter(Q(empresa__isnull=True) | Q(empresa_id__in=ids), activa=True)

    @action(detail=True, methods=["post"])
    def generar_documento(self, request, pk=None):
        """Genera un Documento controlado real desde la plantilla, sustituyendo
        los marcadores {{...}} con datos de la empresa."""
        from datetime import date as _date
        plantilla = self.get_object()
        ids = list(_empresas(request.user))
        emp_id = request.data.get("empresa") or (ids[0] if ids else None)
        if not emp_id or int(emp_id) not in ids:
            return Response({"detail": "Empresa inválida."}, status=400)
        empresa = Empresa.objects.get(id=emp_id)
        try:
            from apps.gestion_documental.models import Documento, TipoDocumento
        except Exception:
            return Response({"detail": "Módulo documental no disponible."}, status=400)

        nombre_emp = empresa.nombre_comercial or empresa.razon_social or ""
        contenido = (plantilla.contenido
                     .replace("{{empresa}}", nombre_emp)
                     .replace("{{razon_social}}", empresa.razon_social or nombre_emp)
                     .replace("{{fecha}}", _date.today().strftime("%d/%m/%Y"))
                     .replace("{{anio}}", str(_date.today().year))
                     .replace("{{codigo}}", plantilla.codigo or "")
                     .replace("{{rfc}}", getattr(empresa, "rfc", "") or ""))
        categoria = plantilla.categoria if plantilla.categoria in (
            "POLITICA", "MANUAL", "PROCEDIMIENTO", "INSTRUCTIVO", "FORMATO", "PLAN") else "OTRO"
        tipo = TipoDocumento.objects.filter(empresa=empresa, categoria=categoria).first()
        if not tipo:
            prefijos = {"POLITICA": "PL", "MANUAL": "MN", "PROCEDIMIENTO": "PR",
                        "INSTRUCTIVO": "IT", "FORMATO": "FR", "PLAN": "PN", "OTRO": "DOC"}
            t_campos = {f.name for f in TipoDocumento._meta.get_fields()}
            t_datos = dict(empresa=empresa, categoria=categoria, nombre=plantilla.get_categoria_display())
            if "codigo" in t_campos:
                t_datos["codigo"] = prefijos.get(categoria, "DOC")
            tipo = TipoDocumento.objects.create(**t_datos)

        campos = {f.name for f in Documento._meta.get_fields()}
        datos = dict(empresa=empresa, tipo=tipo, titulo=plantilla.nombre,
                     descripcion=plantilla.descripcion, estado="BORRADOR")
        if "contenido" in campos:
            datos["contenido"] = contenido
        else:
            datos["descripcion"] = (plantilla.descripcion + "\n\n" + contenido)[:5000]
        if "creado_por" in campos:
            datos["creado_por"] = request.user
        if plantilla.codigo and "codigo" in campos:
            base = plantilla.codigo
            codigo, i = base, 1
            while Documento.objects.filter(empresa=empresa, codigo=codigo).exists():
                i += 1
                codigo = f"{base}-{i}"
            datos["codigo"] = codigo
        doc = Documento.objects.create(**datos)
        return Response({"documento": doc.id, "codigo": getattr(doc, "codigo", ""),
                         "detail": "Documento generado en borrador."}, status=201)

    @action(detail=True, methods=["get"], url_path="word")
    def word(self, request, pk=None):
        """Descarga la plantilla como documento Word (.docx) profesional, con
        encabezado de control documental, control de cambios, firmas y tablas
        con bordes nativas."""
        from datetime import date as _date
        import re as _re
        try:
            from docx import Document  # noqa: F401 — valida disponibilidad de python-docx
        except Exception:
            return Response({"detail": "Generación de Word no disponible en el servidor."}, status=400)
        from django.http import HttpResponse

        plantilla = self.get_object()
        ids = list(_empresas(request.user))
        emp_id = request.query_params.get("empresa") or (ids[0] if ids else None)
        empresa = Empresa.objects.filter(id=emp_id).first() if (emp_id and int(emp_id) in ids) else None
        nombre_emp = (empresa.nombre_comercial or empresa.razon_social) if empresa else "________________"
        rfc = getattr(empresa, "rfc", "") if empresa else ""
        hoy = _date.today()
        fstr = hoy.strftime("%d/%m/%Y")

        def subst(t):
            return (t.replace("{{empresa}}", nombre_emp)
                     .replace("{{razon_social}}", (empresa.razon_social if empresa else nombre_emp))
                     .replace("{{fecha}}", fstr)
                     .replace("{{anio}}", str(hoy.year))
                     .replace("{{codigo}}", plantilla.codigo or "")
                     .replace("{{rfc}}", rfc or ""))

        contenido = subst(plantilla.contenido or "")
        m = _re.search(r"Emisi.n inicial[^\n]*\n", contenido)
        body = contenido[m.end():] if m else contenido

        bio = _docx_controlado(
            empresa=nombre_emp, titulo=plantilla.nombre or "Documento",
            subtitulo=f"{plantilla.get_categoria_display()} · ISO 9001:2015 — Cláusula {plantilla.clausula or '—'}",
            codigo=plantilla.codigo or "—", version="1.0", fecha=fstr, clasificacion="Controlado",
            body=body, elaboro="Responsable de Calidad", reviso="Responsable de Calidad", aprobo="Dirección General")

        fname = (plantilla.codigo or "documento").replace(" ", "_") + ".docx"
        resp = HttpResponse(
            bio.read(),
            content_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document")
        resp["Content-Disposition"] = f'attachment; filename="{fname}"'
        return resp


class HitoCertificacionViewSet(_Scoped):
    queryset = HitoCertificacion.objects.select_related("responsable_user")
    serializer_class = HitoCertificacionSerializer
    filterset_fields = ["empresa", "fase", "completado"]

    def _scoped_ids(self, request):
        ids = list(_empresas(request.user))
        emp = request.query_params.get("empresa")
        return [int(emp)] if emp and int(emp) in ids else ids

    @action(detail=False, methods=["post"])
    def generar(self, request):
        """Crea el roadmap de certificación estándar para la empresa (si no existe)."""
        ids = list(_empresas(request.user))
        emp_id = request.data.get("empresa") or (ids[0] if ids else None)
        if not emp_id or int(emp_id) not in ids:
            return Response({"detail": "Empresa inválida."}, status=400)
        if HitoCertificacion.objects.filter(empresa_id=emp_id).exists():
            return Response({"detail": "El roadmap ya existe.", "creados": 0})
        for orden, (fase, titulo, clausula, ruta) in enumerate(ROADMAP_BASE):
            HitoCertificacion.objects.create(
                empresa_id=emp_id, fase=fase, titulo=titulo, clausula=clausula, ruta=ruta, orden=orden)
        return Response({"creados": len(ROADMAP_BASE)})

    @action(detail=False, methods=["get"])
    def progreso(self, request):
        """Avance del roadmap por fase + siguiente acción recomendada."""
        ids = self._scoped_ids(request)
        hitos = list(HitoCertificacion.objects.filter(empresa_id__in=ids).order_by("orden"))
        fases_def = dict(HitoCertificacion.FASES)
        orden_fases = [f[0] for f in HitoCertificacion.FASES]
        fases = {}
        for h in hitos:
            g = fases.setdefault(h.fase, {"fase": h.fase, "nombre": fases_def.get(h.fase, h.fase),
                                          "total": 0, "completados": 0, "hitos": []})
            g["total"] += 1
            if h.completado:
                g["completados"] += 1
            g["hitos"].append({"id": h.id, "titulo": h.titulo, "clausula": h.clausula,
                               "ruta": h.ruta, "completado": h.completado})
        lista = sorted(fases.values(), key=lambda x: orden_fases.index(x["fase"]) if x["fase"] in orden_fases else 99)
        for g in lista:
            g["progreso"] = round(g["completados"] / g["total"] * 100, 1) if g["total"] else 0
        total = len(hitos)
        hechos = sum(1 for h in hitos if h.completado)
        siguiente = next(({"titulo": h.titulo, "ruta": h.ruta, "fase": fases_def.get(h.fase, h.fase)}
                          for h in hitos if not h.completado), None)
        return Response({
            "fases": lista, "total": total, "completados": hechos,
            "progreso": round(hechos / total * 100, 1) if total else 0,
            "siguiente_accion": siguiente,
        })


class ProgramaAuditoriaViewSet(_Colaborativo):
    queryset = ProgramaAuditoria.objects.select_related("responsable_user")
    serializer_class = ProgramaAuditoriaSerializer
    filterset_fields = ["empresa", "anio", "aprobado"]


class PlantillaChecklistViewSet(viewsets.ModelViewSet):
    permission_classes = [permissions.IsAuthenticated]
    serializer_class = PlantillaChecklistSerializer
    filterset_fields = ["norma", "activa"]

    def get_queryset(self):
        ids = list(_empresas(self.request.user))
        return PlantillaChecklist.objects.filter(Q(empresa__isnull=True) | Q(empresa_id__in=ids)).prefetch_related("items")


class ItemChecklistViewSet(viewsets.ModelViewSet):
    permission_classes = [permissions.IsAuthenticated]
    serializer_class = ItemChecklistSerializer
    filterset_fields = ["checklist"]

    def get_queryset(self):
        ids = list(_empresas(self.request.user))
        return ItemChecklist.objects.filter(Q(checklist__empresa__isnull=True) | Q(checklist__empresa_id__in=ids))


class RegistroCalidadViewSet(_Colaborativo):
    queryset = RegistroCalidad.objects.select_related("responsable_user", "proceso_ref")
    serializer_class = RegistroCalidadSerializer
    filterset_fields = ["empresa", "soporte", "disposicion", "activo", "responsable_user"]


class GestionCambioViewSet(_Colaborativo):
    queryset = GestionCambio.objects.select_related("responsable_user", "creado_por")
    serializer_class = GestionCambioSerializer
    filterset_fields = ["empresa", "estado", "responsable_user", "tipo", "prioridad"]


class ComunicacionSGCViewSet(_Colaborativo):
    queryset = ComunicacionSGC.objects.select_related("responsable_user")
    serializer_class = ComunicacionSGCSerializer
    filterset_fields = ["empresa", "direccion", "activo", "responsable_user"]


class ConocimientoOrganizacionalViewSet(_Scoped):
    queryset = ConocimientoOrganizacional.objects.select_related("autor_user")
    serializer_class = ConocimientoOrganizacionalSerializer
    filterset_fields = ["empresa", "tipo"]
    search_fields = ["titulo", "contenido", "etiquetas", "area"]

    def perform_create(self, serializer):
        serializer.save(autor_user=self.request.user)


# Procesos, objetivos y riesgos típicos por giro de empresa (onboarding).
GIROS_SGC = {
    "manufactura": {
        "label": "Manufactura / Producción",
        "procesos": [
            ("ESTRATEGICO", "Planeación estratégica"), ("CLAVE", "Producción"),
            ("CLAVE", "Control de calidad"), ("CLAVE", "Compras"),
            ("APOYO", "Mantenimiento"), ("APOYO", "Recursos humanos")],
        "objetivos": [("Reducir el % de producto no conforme", "≤ 2%"),
                      ("Cumplir el programa de producción", "≥ 95%"),
                      ("Mantener la satisfacción del cliente", "≥ 90%")],
        "riesgos": [("Falla de equipo crítico en producción", 3, 4),
                    ("Desabasto de materia prima", 3, 4),
                    ("Producto fuera de especificación", 2, 5)],
    },
    "servicios": {
        "label": "Servicios",
        "procesos": [
            ("ESTRATEGICO", "Dirección"), ("CLAVE", "Prestación del servicio"),
            ("CLAVE", "Atención al cliente"), ("CLAVE", "Ventas"),
            ("APOYO", "Recursos humanos"), ("APOYO", "Tecnología")],
        "objetivos": [("Mejorar el tiempo de respuesta al cliente", "≤ 24 h"),
                      ("Aumentar la satisfacción del cliente (CSAT)", "≥ 90%"),
                      ("Reducir quejas recurrentes", "≤ 5/mes")],
        "riesgos": [("Insatisfacción del cliente por demoras", 3, 4),
                    ("Rotación de personal clave", 3, 3),
                    ("Incumplimiento de acuerdos de servicio (SLA)", 2, 4)],
    },
    "comercio": {
        "label": "Comercio / Distribución",
        "procesos": [
            ("ESTRATEGICO", "Dirección comercial"), ("CLAVE", "Compras"),
            ("CLAVE", "Almacén y distribución"), ("CLAVE", "Ventas"),
            ("APOYO", "Administración"), ("APOYO", "Recursos humanos")],
        "objetivos": [("Cumplir las entregas a tiempo", "≥ 95%"),
                      ("Reducir faltantes de inventario", "≤ 3%"),
                      ("Aumentar la satisfacción del cliente", "≥ 90%")],
        "riesgos": [("Entregas fuera de tiempo", 3, 4),
                    ("Diferencias de inventario", 3, 3),
                    ("Proveedores no confiables", 2, 4)],
    },
    "transporte": {
        "label": "Transporte / Logística",
        "procesos": [
            ("ESTRATEGICO", "Dirección"), ("CLAVE", "Operación de transporte"),
            ("CLAVE", "Mantenimiento de flota"), ("CLAVE", "Atención al cliente"),
            ("APOYO", "Recursos humanos"), ("APOYO", "Compras")],
        "objetivos": [("Cumplir las entregas a tiempo", "≥ 95%"),
                      ("Reducir incidentes/accidentes", "0 graves"),
                      ("Disponibilidad de la flota", "≥ 90%")],
        "riesgos": [("Accidente de unidad", 2, 5),
                    ("Retraso en entregas", 3, 4),
                    ("Falla mecánica de flota", 3, 3)],
    },
    "general": {
        "label": "General / Otro",
        "procesos": [
            ("ESTRATEGICO", "Dirección"), ("CLAVE", "Operación principal"),
            ("CLAVE", "Atención al cliente"), ("APOYO", "Compras"),
            ("APOYO", "Recursos humanos")],
        "objetivos": [("Aumentar la satisfacción del cliente", "≥ 90%"),
                      ("Cumplir los objetivos operativos", "≥ 95%"),
                      ("Reducir no conformidades", "tendencia a la baja")],
        "riesgos": [("Insatisfacción del cliente", 3, 4),
                    ("Incumplimiento de requisitos legales", 2, 5),
                    ("Falta de competencia del personal", 2, 3)],
    },
}


class OnboardingSGCViewSet(viewsets.ViewSet):
    """Asistente de alta del SGC: genera la base (procesos, objetivos, riesgos,
    roadmap) según el giro, para que una empresa nueva arranque en minutos."""
    permission_classes = [permissions.IsAuthenticated]

    @action(detail=False, methods=["get"])
    def giros(self, request):
        return Response([{"id": k, "label": v["label"]} for k, v in GIROS_SGC.items()])

    @action(detail=False, methods=["get"])
    def estado(self, request):
        """Indica si la empresa ya tiene contenido del SGC."""
        ids = list(_empresas(request.user))
        emp = request.query_params.get("empresa")
        emp_id = int(emp) if emp and int(emp) in ids else (ids[0] if ids else None)
        if not emp_id:
            return Response({"detail": "Empresa inválida."}, status=400)
        return Response({
            "procesos": Proceso.objects.filter(empresa_id=emp_id).count(),
            "objetivos": ObjetivoCalidad.objects.filter(empresa_id=emp_id).count(),
            "riesgos": Riesgo.objects.filter(empresa_id=emp_id).count(),
            "roadmap": HitoCertificacion.objects.filter(empresa_id=emp_id).count(),
            "diagnostico": EvaluacionRequisito.objects.filter(empresa_id=emp_id).count(),
        })

    @action(detail=False, methods=["post"])
    def generar(self, request):
        """Crea procesos SIPOC, objetivos, riesgos y roadmap base según el giro.
        No duplica lo que ya exista."""
        ids = list(_empresas(request.user))
        emp_id = request.data.get("empresa") or (ids[0] if ids else None)
        if not emp_id or int(emp_id) not in ids:
            return Response({"detail": "Empresa inválida."}, status=400)
        giro = request.data.get("giro", "general")
        cfg = GIROS_SGC.get(giro, GIROS_SGC["general"])
        from datetime import date as _date, timedelta as _td
        res = {"procesos": 0, "objetivos": 0, "riesgos": 0, "roadmap": 0}

        for i, (tipo, nombre) in enumerate(cfg["procesos"]):
            obj, created = Proceso.objects.get_or_create(
                empresa_id=emp_id, nombre=nombre,
                defaults={"tipo": tipo, "codigo": f"P-{i + 1:02d}"})
            if created:
                res["procesos"] += 1

        for objetivo, meta in cfg["objetivos"]:
            _, created = ObjetivoCalidad.objects.get_or_create(
                empresa_id=emp_id, objetivo=objetivo,
                defaults={"meta": meta, "estado": "EN_CURSO",
                          "fecha_limite": _date.today() + _td(days=365)})
            if created:
                res["objetivos"] += 1

        for descripcion, prob, imp in cfg["riesgos"]:
            _, created = Riesgo.objects.get_or_create(
                empresa_id=emp_id, descripcion=descripcion,
                defaults={"probabilidad": prob, "impacto": imp, "estado": "IDENTIFICADO"})
            if created:
                res["riesgos"] += 1

        if not HitoCertificacion.objects.filter(empresa_id=emp_id).exists():
            for orden, (fase, titulo, clausula, ruta) in enumerate(ROADMAP_BASE):
                HitoCertificacion.objects.create(
                    empresa_id=emp_id, fase=fase, titulo=titulo, clausula=clausula,
                    ruta=ruta, orden=orden)
            res["roadmap"] = len(ROADMAP_BASE)

        return Response({"giro": cfg["label"], "creados": res})
