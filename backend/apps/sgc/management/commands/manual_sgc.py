# -*- coding: utf-8 -*-
"""Genera un Manual de Uso del SGC completo en Word (.docx).

Cubre todos los módulos del sistema organizados por el ciclo PHVA, con para qué
sirve cada uno (cláusula ISO 9001), por qué importa y cómo usarlo paso a paso.

Uso:
    python manage.py manual_sgc                 # guarda en el Escritorio
    python manage.py manual_sgc --salida ruta.docx
"""
from __future__ import annotations

import os
from datetime import date

from django.core.management.base import BaseCommand


# ── Helpers de formato docx ───────────────────────────────────────────────────
def _setup(doc):
    from docx.shared import Pt
    sec = doc.sections[0]
    from docx.shared import Inches
    sec.left_margin = sec.right_margin = Inches(0.9)
    sec.top_margin = Inches(0.8); sec.bottom_margin = Inches(0.8)
    n = doc.styles["Normal"]; n.font.name = "Calibri"; n.font.size = Pt(10.5)


def _rule(p, color="4F46E5", sz=6):
    from docx.oxml.ns import qn
    from docx.oxml import OxmlElement
    pPr = p._p.get_or_add_pPr()
    pbdr = OxmlElement("w:pBdr"); b = OxmlElement("w:bottom")
    b.set(qn("w:val"), "single"); b.set(qn("w:sz"), str(sz)); b.set(qn("w:space"), "3"); b.set(qn("w:color"), color)
    pbdr.append(b); pPr.append(pbdr)


def _field(par, instr):
    from docx.oxml.ns import qn
    from docx.oxml import OxmlElement
    run = par.add_run()
    for t, v in (("begin", None), (None, instr), ("end", None)):
        if t:
            e = OxmlElement("w:fldChar"); e.set(qn("w:fldCharType"), t); run._r.append(e)
        else:
            it = OxmlElement("w:instrText"); it.set(qn("xml:space"), "preserve"); it.text = v; run._r.append(it)


class Command(BaseCommand):
    help = "Genera el Manual de Uso del SGC en Word (.docx)."

    def add_arguments(self, parser):
        parser.add_argument("--salida", type=str, default=None)

    def handle(self, *args, **o):
        try:
            from docx import Document
            from docx.shared import Pt, RGBColor
            from docx.enum.text import WD_ALIGN_PARAGRAPH
        except Exception:
            self.stderr.write(self.style.ERROR("Falta python-docx. Instala con: pip install python-docx")); return

        INDIGO = RGBColor(0x4F, 0x46, 0xE5)
        DARK = RGBColor(0x1E, 0x29, 0x3B)
        GREY = RGBColor(0x64, 0x74, 0x8B)
        CENTER = WD_ALIGN_PARAGRAPH.CENTER

        doc = Document()
        _setup(doc)

        def h1(txt):
            doc.add_page_break()
            p = doc.add_paragraph(); r = p.add_run(txt); r.bold = True; r.font.size = Pt(18); r.font.color.rgb = INDIGO
            _rule(p, sz=10)
            return p

        def h2(txt):
            p = doc.add_paragraph(); p.paragraph_format.space_before = Pt(10)
            r = p.add_run(txt); r.bold = True; r.font.size = Pt(13.5); r.font.color.rgb = DARK
            _rule(p)
            return p

        def h3(txt):
            p = doc.add_paragraph(); p.paragraph_format.space_before = Pt(6)
            r = p.add_run(txt); r.bold = True; r.font.size = Pt(11.5); r.font.color.rgb = INDIGO
            return p

        def para(txt, size=10.5):
            p = doc.add_paragraph(); r = p.add_run(txt); r.font.size = Pt(size); return p

        def bullets(items):
            for it in items:
                p = doc.add_paragraph(style="List Bullet"); r = p.add_run(it); r.font.size = Pt(10.5)

        def steps(items):
            for i, it in enumerate(items, 1):
                p = doc.add_paragraph(style="List Number"); r = p.add_run(it); r.font.size = Pt(10.5)

        def modulo(nombre, clausula, ruta, para_que, por_que, como):
            h3(f"{nombre}")
            sub = doc.add_paragraph()
            r = sub.add_run(f"ISO 9001 · {clausula}    |    Ruta: {ruta}")
            r.italic = True; r.font.size = Pt(8.5); r.font.color.rgb = GREY
            p = doc.add_paragraph(); r = p.add_run("Para qué sirve: "); r.bold = True; r.font.size = Pt(10)
            r2 = p.add_run(para_que); r2.font.size = Pt(10)
            p = doc.add_paragraph(); r = p.add_run("Por qué importa: "); r.bold = True; r.font.size = Pt(10)
            r2 = p.add_run(por_que); r2.font.size = Pt(10)
            cp = doc.add_paragraph(); cr = cp.add_run("Cómo usarlo:"); cr.bold = True; cr.font.size = Pt(10)
            steps(como)

        # ───────────────────────── PORTADA ─────────────────────────
        for _ in range(4):
            doc.add_paragraph("")
        p = doc.add_paragraph(); p.alignment = CENTER
        r = p.add_run("MANUAL DE USO"); r.bold = True; r.font.size = Pt(34); r.font.color.rgb = INDIGO
        p = doc.add_paragraph(); p.alignment = CENTER
        r = p.add_run("Sistema de Gestión de Calidad"); r.bold = True; r.font.size = Pt(22); r.font.color.rgb = DARK
        p = doc.add_paragraph(); p.alignment = CENTER
        r = p.add_run("ISO 9001:2015"); r.font.size = Pt(16); r.font.color.rgb = GREY
        doc.add_paragraph("")
        p = doc.add_paragraph(); p.alignment = CENTER
        r = p.add_run("Guía completa de los módulos, su propósito y su uso paso a paso"); r.italic = True; r.font.size = Pt(11); r.font.color.rgb = GREY
        for _ in range(8):
            doc.add_paragraph("")
        p = doc.add_paragraph(); p.alignment = CENTER
        r = p.add_run(f"Versión 1.0  ·  {date.today().strftime('%d/%m/%Y')}"); r.font.size = Pt(10); r.font.color.rgb = GREY

        # Pie con número de página.
        fp = doc.sections[0].footer.paragraphs[0]; fp.alignment = CENTER
        r = fp.add_run("Manual de Uso del SGC · ISO 9001:2015    —    Página "); r.font.size = Pt(8); r.font.color.rgb = GREY
        _field(fp, "PAGE")
        r = fp.add_run(" de "); r.font.size = Pt(8); r.font.color.rgb = GREY
        _field(fp, "NUMPAGES")

        # ───────────────────────── 1. INTRODUCCIÓN ─────────────────────────
        h1("1. Introducción")
        para("Este manual explica cómo usar el módulo de Calidad (SGC) del ERP para implementar y "
             "operar un Sistema de Gestión de Calidad conforme a la norma ISO 9001:2015. Está dirigido "
             "a responsables de calidad, dueños de proceso y a todo el equipo que participa en el sistema.")
        h2("¿Qué es un SGC?")
        para("Un Sistema de Gestión de Calidad es el conjunto de procesos, documentos, responsabilidades "
             "y controles con los que una organización asegura que sus productos y servicios cumplen los "
             "requisitos del cliente, los legales y los propios, mejorando de forma continua. ISO 9001:2015 "
             "es la norma internacional que define los requisitos para lograrlo.")
        h2("Principios sobre los que trabaja el sistema")
        bullets([
            "Enfoque a procesos: la organización opera como una red de procesos con entradas, salidas e indicadores.",
            "Ciclo PHVA (Planear–Hacer–Verificar–Actuar): toda mejora gira sobre este ciclo. Los módulos están organizados según esta lógica.",
            "Pensamiento basado en riesgos: se identifican y tratan riesgos y oportunidades antes de que ocurran problemas.",
            "Mejora continua: la información de auditorías, quejas, KPIs y no conformidades alimenta acciones de mejora.",
        ])

        # ───────────────────────── 2. PRIMEROS PASOS ─────────────────────────
        h1("2. Primeros pasos")
        h2("Acceso y empresa activa")
        steps([
            "Inicia sesión con tu usuario y contraseña.",
            "Selecciona la empresa activa en el selector superior. Todo lo que veas y registres pertenece a esa empresa.",
            "Entra al módulo de Calidad (SGC) desde el menú principal para llegar al Centro de Mando.",
        ])
        h2("El Centro de Mando (Dashboard del SGC)")
        para("La página principal del SGC (/sgc) muestra el estado global del sistema:")
        bullets([
            "Madurez del SGC: índice de 0 a 100 % que promedia 6 dimensiones (cumplimiento ISO, implementación, "
            "objetivos, KPIs en meta, NC resueltas y riesgos controlados). Indica qué tan operando y vivo está el sistema.",
            "Cumplimiento ISO: porcentaje del diagnóstico/gap analysis. Mide cuánto de la norma está cubierto.",
            "Importante: puedes tener alto cumplimiento documental y baja madurez. Significa que el sistema está "
            "bien diseñado pero todavía no se ha implementado ni operado en el día a día.",
            "Evolución y alertas: gráfica de la madurez en el tiempo, proyección a la certificación y alertas "
            "inteligentes (NC por vencer, riesgos críticos, calibraciones vencidas, KPIs fuera de meta).",
            "Rueda PHVA y módulos: navega a cada módulo desde las tarjetas, agrupadas por fase del ciclo.",
        ])
        h2("Cómo se mide el avance hacia la certificación")
        para("El sistema separa dos ideas que conviene no confundir: 'saber qué pide la norma' (cumplimiento ISO, "
             "que sale del Diagnóstico) y 'tener el sistema funcionando' (madurez, que mide evidencia real de "
             "ejecución). Para certificarte necesitas ambas: documentar el SGC y demostrar que opera con eficacia.")

        # ───────────────────────── 3. PLANEAR ─────────────────────────
        h1("3. PLANEAR — Contexto, liderazgo y planificación")
        para("Fase donde se define el rumbo del SGC: el contexto, las partes interesadas, la política, los "
             "objetivos, los riesgos y el plan para implementar la norma. Cláusulas ISO 4, 5 y 6.")

        modulo("Contexto y Partes Interesadas", "4.1 y 4.2", "/sgc/contexto",
               "Determinar las cuestiones internas y externas (análisis FODA) y las partes interesadas con sus necesidades.",
               "La norma exige entender el entorno y a quién debe responder la organización; de aquí salen riesgos y oportunidades.",
               ["Registra fortalezas, debilidades, oportunidades y amenazas (FODA) con su estrategia asociada.",
                "Da de alta cada parte interesada (clientes, proveedores, personal, autoridades…) con sus necesidades e influencia.",
                "Revisa esta información al menos una vez al año o cuando cambie el contexto."])

        modulo("Política y Objetivos de Calidad", "5.2 y 6.2", "/sgc/politica",
               "Declarar el compromiso de la dirección (política) y fijar objetivos de calidad medibles.",
               "La política da el marco; los objetivos convierten la intención en metas concretas con seguimiento.",
               ["Redacta o ajusta la política de calidad y publícala.",
                "Define objetivos medibles con su indicador, meta, responsable y plazo.",
                "Actualiza el avance de cada objetivo; alimenta la madurez del sistema."])

        modulo("Gestión de Riesgos", "6.1", "/sgc/riesgos",
               "Identificar, evaluar y tratar riesgos y oportunidades por proceso (matriz probabilidad × impacto).",
               "Anticiparse a los problemas es el corazón del pensamiento basado en riesgos que pide ISO 9001.",
               ["Registra cada riesgo u oportunidad y asígnale un proceso y responsable.",
                "Evalúa probabilidad e impacto (1–5); el sistema calcula la severidad.",
                "Define el plan de tratamiento (evitar, reducir, transferir, aceptar) y cambia el estado a Controlado cuando proceda."])

        modulo("Diagnóstico ISO (Gap Analysis)", "6.1", "/sgc/diagnostico",
               "Evaluar requisito por requisito el cumplimiento de la norma (Cumple / Parcial / No / N/A).",
               "Es el punto de partida: muestra el % de cumplimiento, detecta brechas y genera el plan de acción.",
               ["Recorre el cuestionario y marca el nivel de cada requisito.",
                "Para las brechas (No / Parcial), escribe la observación o acción para cerrarlas.",
                "Usa 'Generar plan de acción (CAPA)' para crear no conformidades automáticamente desde las brechas."])

        modulo("Gestión de Cambios", "6.3", "/sgc/cambios",
               "Planificar los cambios del SGC de forma controlada: propósito, consecuencias, integridad, recursos y responsables.",
               "Cambiar 'sobre la marcha' puede romper procesos o perder trazabilidad; la norma exige planearlos.",
               ["Crea el cambio con el asistente por pasos: identificación, análisis 6.3, riesgo y recursos, plan y responsables.",
                "Evalúa el riesgo del cambio en la matriz 5×5 y desglosa el plan de acción en tareas.",
                "Avanza el cambio por el flujo: Propuesto → Aprobado → En proceso → Implementado."])

        modulo("Tablero de Implementación", "6.2", "/sgc/implementacion",
               "Convertir las cláusulas de la norma en tareas accionables y repartirlas en el equipo (Kanban).",
               "Certificarse no es solo documentar: hay que implementar cada requisito; el tablero asegura que nada quede sin hacer.",
               ["Pulsa 'Generar desde ISO' para crear una tarjeta por cada cláusula.",
                "Asigna responsable, prioridad y fecha límite; desglosa cada tarjeta en subtareas con peso.",
                "Arrastra las tarjetas Por hacer → En proceso → Implementado → Verificado conforme avanzan."])

        # ───────────────────────── 4. HACER ─────────────────────────
        h1("4. HACER — Apoyo y operación")
        para("Fase de ejecución: recursos, personas, documentos, comunicación y la operación de los procesos. "
             "Cláusulas ISO 7 y 8.")

        modulo("Biblioteca de Plantillas", "7.5", "/sgc/plantillas",
               "Generar documentos del SGC a partir de plantillas ISO listas para usar.",
               "Redactar desde cero es el mayor freno para certificarse; las plantillas aceleran la documentación.",
               ["Filtra por fase PHVA o categoría y usa 'Vista' para previsualizar.",
                "Pulsa 'Generar' para crear un documento controlado en borrador con los datos de tu empresa.",
                "Descarga cualquier plantilla o documento como Word (.docx) con tablas y formato profesional."])

        modulo("Gestión Documental", "7.5", "/documentos",
               "Administrar los documentos controlados: código, versión, estado, aprobación y distribución.",
               "La información documentada debe estar controlada: versiones vigentes, aprobaciones y trazabilidad.",
               ["Crea un documento subiendo un archivo o redactando su contenido con el editor tipo Word.",
                "Envíalo a aprobación; al aprobarse pasa a Vigente. Las versiones quedan en el historial.",
                "Visualízalo con zoom, imprímelo o descárgalo como Word desde su ficha."])

        modulo("Lista Maestra de Registros", "7.5.3", "/sgc/registros",
               "Controlar los registros: soporte, ubicación, tiempo de retención y disposición.",
               "Los registros son la evidencia del SGC; deben conservarse y eliminarse de forma controlada.",
               ["Da de alta cada registro con su código, proceso y responsable.",
                "Define soporte (digital/físico), ubicación, meses de retención y disposición final.",
                "Mantén la lista al día para auditorías."])

        modulo("Matriz de Comunicación", "7.4", "/sgc/comunicacion",
               "Definir qué se comunica, cuándo, a quién, cómo y quién es responsable.",
               "La norma exige comunicaciones internas y externas pertinentes y planificadas.",
               ["Registra cada comunicación con su dirección (interna/externa) y canal.",
                "Asigna responsable y frecuencia.",
                "Revisa la matriz cuando cambien procesos o partes interesadas."])

        modulo("Base de Conocimiento", "7.1.6", "/sgc/conocimiento",
               "Preservar lecciones aprendidas, mejores prácticas y fuentes externas.",
               "Evita repetir errores y que el conocimiento se pierda cuando alguien deja la organización.",
               ["Captura el conocimiento al cerrar proyectos, auditorías o incidentes.",
                "Clasifícalo por tipo y área, y agrégale etiquetas para encontrarlo.",
                "Consúltalo antes de iniciar trabajos similares."])

        modulo("Capacitación", "7.2", "/sgc/capacitacion",
               "Programar cursos, registrar asistencia, evidencias y evaluar la eficacia de la formación.",
               "El personal debe ser competente; la capacitación cierra las brechas detectadas.",
               ["Programa cada curso con responsable y participantes.",
                "Registra la realización y la evidencia.",
                "Evalúa la eficacia tras el curso."])

        modulo("Competencias", "7.2", "/sgc/competencias",
               "Definir perfiles de puesto, evaluar al personal y detectar brechas de competencia.",
               "Asegura que cada persona tiene el nivel requerido para que su trabajo no afecte la calidad.",
               ["Define el perfil de cada puesto con sus competencias y nivel requerido.",
                "Evalúa a cada persona; el sistema calcula la brecha y muestra la matriz y el radar.",
                "Genera el plan de capacitación desde la brecha y verifica la eficacia (cierre del ciclo)."])

        modulo("Control de Equipos", "7.1.5", "/sgc/equipos",
               "Inventariar equipos de seguimiento y medición y controlar sus calibraciones.",
               "Mediciones fiables requieren equipos calibrados y trazables; el sistema alerta de vencimientos.",
               ["Registra cada equipo y su frecuencia de calibración.",
                "Programa y registra las calibraciones con su certificado.",
                "Atiende las alertas de calibración vencida."])

        modulo("Mapa de Procesos (SIPOC)", "8.1", "/sgc/procesos",
               "Caracterizar cada proceso: dueño, entradas/salidas, proveedores, clientes, KPIs y riesgos.",
               "El enfoque a procesos exige conocer cómo se conectan y miden los procesos de la organización.",
               ["Da de alta cada proceso con su dueño y tipo (estratégico/clave/apoyo).",
                "Completa el SIPOC y vincula indicadores y riesgos.",
                "Úsalo como base para auditorías y mejora."])

        modulo("Evaluación de Proveedores", "8.4", "/sgc/evaluacion-proveedores",
               "Seleccionar, evaluar y dar seguimiento a los proveedores externos.",
               "Los proveedores afectan la calidad final; deben controlarse según su impacto.",
               ["Evalúa con los criterios definidos (calidad, entrega, servicio, documentación).",
                "Clasifica al proveedor y define plan de mejora si es necesario.",
                "Reevalúa periódicamente."])

        modulo("Salidas No Conformes", "8.7", "/sgc/salidas-no-conformes",
               "Identificar y controlar producto/servicio que no cumple, y su disposición.",
               "Evita el uso o entrega no intencionada de salidas no conformes.",
               ["Registra la salida no conforme y su origen.",
                "Define la disposición (corrección, segregación, concesión, devolución, desecho…).",
                "Verifica la conformidad tras la corrección y conserva la evidencia."])

        # ───────────────────────── 5. VERIFICAR ─────────────────────────
        h1("5. VERIFICAR — Evaluación del desempeño")
        para("Fase de medición y seguimiento: indicadores, satisfacción del cliente, auditorías y revisión "
             "por la dirección. Cláusula ISO 9.")

        modulo("Indicadores KPI", "9.1", "/sgc/kpis",
               "Medir el desempeño de los procesos con metas, mediciones periódicas y tendencias.",
               "Lo que no se mide no se mejora; los KPIs muestran si el SGC logra sus objetivos.",
               ["Define cada KPI con su meta y sentido (mayor o menor es mejor).",
                "Captura mediciones periódicas; el sistema dibuja la tendencia y el semáforo.",
                "Atiende los KPIs fuera de meta (rojo/amarillo)."])

        modulo("Quejas y Satisfacción", "9.1.2", "/sgc/quejas",
               "Registrar quejas, felicitaciones y medir la satisfacción del cliente (CSAT/NPS).",
               "La percepción del cliente es una entrada clave del desempeño del SGC.",
               ["Registra cada caso con su tipo y calificación de satisfacción.",
                "Da seguimiento hasta cerrarlo.",
                "Analiza las estadísticas de satisfacción."])

        modulo("Encuestas de Cliente", "9.1.2", "/sgc/encuestas",
               "Crear encuestas con liga pública para recoger la voz del cliente.",
               "Automatiza la captura de satisfacción; las respuestas llegan al módulo de Quejas.",
               ["Crea la encuesta y comparte su liga pública.",
                "Las respuestas se tabulan automáticamente.",
                "Usa los resultados en la revisión por la dirección."])

        modulo("Programa de Auditorías", "9.2", "/sgc/programa-auditorias",
               "Planificar el programa anual de auditorías internas por proceso.",
               "Las auditorías deben planificarse: frecuencia, métodos y responsabilidades.",
               ["Crea el programa del año con objetivo, alcance y criterios.",
                "Asigna procesos y meses a auditar.",
                "Apruébalo para que sirva de base a las auditorías."])

        modulo("Auditorías Internas", "9.2", "/sgc/auditorias",
               "Ejecutar auditorías con checklist en vivo y generar hallazgos.",
               "Verifican si el SGC es conforme y eficaz, y alimentan la mejora.",
               ["Crea la auditoría y abre su checklist en vivo.",
                "Marca cada requisito (Conforme / No conforme / Observación) con notas.",
                "Cierra la auditoría: se generan automáticamente los hallazgos y no conformidades."])

        modulo("Revisión por la Dirección", "9.3", "/sgc/revision-direccion",
               "Registrar el acta de la revisión del SGC por la alta dirección.",
               "La dirección debe revisar el sistema periódicamente para asegurar su conveniencia y eficacia.",
               ["Registra la reunión con sus participantes y periodo.",
                "Documenta las entradas (9.3.2) y las salidas/acuerdos (9.3.3).",
                "Da seguimiento a los acuerdos."])

        # ───────────────────────── 6. ACTUAR ─────────────────────────
        h1("6. ACTUAR — Mejora")
        para("Fase de mejora: tratar las no conformidades, aplicar acciones correctivas y mantener el sistema "
             "vivo. Cláusula ISO 10.")

        modulo("No Conformidades (CAPA)", "10.2", "/sgc/no-conformidades",
               "Gestionar no conformidades con análisis de causa raíz, plan de acción y verificación de eficacia.",
               "Eliminar la causa de los problemas evita que se repitan; es el motor de la mejora continua.",
               ["Registra la no conformidad con su origen y descripción.",
                "Analiza la causa raíz (5 Porqués e Ishikawa 6M) y clasifícala.",
                "Define el plan de acción correctiva, impleméntalo y verifica su eficacia antes de cerrar.",
                "Consulta la sección de Análisis (Pareto de causas, tendencia mensual, origen) para ver patrones."])

        modulo("Agenda de Cumplimiento", "10.3", "/sgc/agenda",
               "Ver en un solo lugar todos los vencimientos del SGC para mantenerlo vivo.",
               "Un SGC se mantiene con disciplina; la agenda evita que se venzan calibraciones, auditorías o acciones.",
               ["Revisa la agenda periódicamente.",
                "Atiende los vencimientos próximos.",
                "Úsala como checklist de mantenimiento del sistema."])

        # ───────────────────────── 7. CAMINO A LA CERTIFICACIÓN ─────────────────────────
        h1("7. Flujo recomendado hacia la certificación")
        steps([
            "Diagnóstico inicial: completa el gap analysis para conocer tu punto de partida.",
            "Planeación y liderazgo: define contexto, partes interesadas, política, objetivos y riesgos.",
            "Documentación: genera tu documentación base con las plantillas y contrólala en Gestión Documental.",
            "Implementación: usa el tablero para implementar cada cláusula en equipo.",
            "Medición: define KPIs, mide la satisfacción del cliente y registra evidencia.",
            "Auditoría interna: ejecuta auditorías con el checklist en vivo y registra hallazgos.",
            "Revisión por la dirección: documenta el acta con entradas y salidas.",
            "Acciones y pre-auditoría: cierra las no conformidades con CAPA y verifica su eficacia.",
            "Auditoría de certificación: con la madurez alta y la evidencia lista, enfrenta la auditoría externa.",
        ])
        para("Sigue el avance global en el Centro de Mando: cuando la madurez se acerca al 85 % y las brechas "
             "están cerradas, el sistema está listo para la certificación.")

        # ───────────────────────── 8. COLABORACIÓN Y FAQ ─────────────────────────
        h1("8. Colaboración, roles y preguntas frecuentes")
        h2("Colaboración")
        bullets([
            "Cada registro puede asignarse a un responsable, que recibe una notificación (campana).",
            "Los módulos llevan bitácora de actividad y permiten comentarios entre el equipo.",
            "La 'carga del equipo' en el dashboard muestra quién tiene más pendientes abiertos.",
        ])
        h2("Preguntas frecuentes")
        h3("¿Por qué mi cumplimiento ISO es alto pero la madurez es baja?")
        para("Porque la madurez mide la operación real (implementación, objetivos, KPIs, cierre de NC, riesgos), "
             "no solo lo documentado. Es normal al inicio: el sistema está diseñado pero aún no operando.")
        h3("¿Cómo subo la madurez?")
        para("Avanza el tablero de implementación a 'Verificado', actualiza el avance de los objetivos, cumple los "
             "KPIs, cierra no conformidades y controla los riesgos.")
        h3("¿Puedo trabajar con varias empresas?")
        para("Sí. Cada empresa tiene su propio SGC; cambia la empresa activa en el selector superior.")
        h3("¿Cómo obtengo mis documentos en Word?")
        para("Desde la Biblioteca de Plantillas o desde la ficha de cualquier documento, usa el botón 'Word' / "
             "'Descargar Word'. El archivo se genera con encabezado de control, tablas y firmas.")

        # ───────────────────────── 9. GLOSARIO ─────────────────────────
        h1("9. Glosario")
        glos = [
            ("SGC", "Sistema de Gestión de Calidad."),
            ("PHVA", "Planear, Hacer, Verificar, Actuar: ciclo de mejora continua."),
            ("Gap analysis", "Diagnóstico que compara el estado actual contra los requisitos de la norma."),
            ("No conformidad (NC)", "Incumplimiento de un requisito."),
            ("CAPA", "Acción Correctiva y Preventiva."),
            ("Causa raíz", "Origen real de un problema; se analiza con 5 Porqués o Ishikawa."),
            ("KPI", "Indicador clave de desempeño."),
            ("Madurez", "Índice que mide qué tan operando está el SGC (6 dimensiones)."),
            ("Parte interesada", "Persona u organización que afecta o es afectada por el SGC."),
            ("Información documentada", "Documentos y registros controlados del sistema."),
        ]
        for term, d in glos:
            p = doc.add_paragraph()
            r = p.add_run(f"{term}: "); r.bold = True; r.font.size = Pt(10.5)
            r2 = p.add_run(d); r2.font.size = Pt(10.5)

        # ───────────────────────── Guardar ─────────────────────────
        salida = o.get("salida")
        if not salida:
            escritorio = os.path.join(os.path.expanduser("~"), "OneDrive", "Desktop")
            if not os.path.isdir(escritorio):
                escritorio = os.path.join(os.path.expanduser("~"), "Desktop")
            if not os.path.isdir(escritorio):
                escritorio = os.path.expanduser("~")
            salida = os.path.join(escritorio, "Manual_de_Uso_SGC_ISO9001.docx")
        doc.save(salida)
        self.stdout.write(self.style.SUCCESS(f"Manual generado: {salida}"))
